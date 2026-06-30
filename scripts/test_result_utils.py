import json
import re
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import load_workbook

from utils import EXCEL_HEADERS


KNOWN_BAD_DISEASE_TERMS = [
    "澄迈县的李女士去年帕金森病",
    "旨在减轻被保险人因患大病",
    "特药范围以及覆盖的疾病病",
    "了解症",
    "妥妥的花小钱保大病",
    "医保定点医药机构发生的大病",
    "中国居民营养与慢性病",
    "合理自费费用",
    "高额医疗费用",
    "参保群众",
    "美国国家",
    "加拿大国立",
    "英国癌",
    "欧洲癌",
    "高价自费",
    "一旦患",
]
GENERIC_DISEASE_TERMS = {"大病", "既往症", "慢性病", "特殊病"}


def desktop_result_dir(name: str = "DataExtractor_test_results") -> Path:
    root = Path.home() / "Desktop" / f"{name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    root.mkdir(parents=True, exist_ok=True)
    return root


def run_command(command: List[str], cwd: Path, log_path: Path, timeout: int = 1800) -> Dict[str, Any]:
    log_path.parent.mkdir(parents=True, exist_ok=True)
    completed = subprocess.run(command, cwd=cwd, text=True, capture_output=True, timeout=timeout)
    output = (completed.stdout or "") + ("\n" + completed.stderr if completed.stderr else "")
    log_path.write_text(mask_runtime_output(output), encoding="utf-8")
    return {"command": command, "returncode": completed.returncode, "log_path": str(log_path)}


def mask_runtime_output(text: str) -> str:
    text = re.sub(r"sk-[A-Za-z0-9_-]+", "sk-***", text or "")
    text = re.sub(r"(mysql\+pymysql://[^:/@\s]+:)([^@\s]+)(@)", r"\1***\3", text)
    text = re.sub(r"(?i)(password|token|secret|api_key)\s*[=:]\s*([^\s,;]+)", r"\1=***", text)
    return text


def latest_xlsx(directory: Path) -> Path:
    files = sorted(directory.glob("*.xlsx"), key=lambda item: item.stat().st_mtime)
    if not files:
        raise FileNotFoundError(f"未找到 Excel 输出: {directory}")
    return files[-1]


def validate_output_workbook(path: Path) -> Dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_names = workbook.sheetnames
        result_sheet = workbook["结果数据"] if "结果数据" in sheet_names else workbook[sheet_names[0]]
        rows = list(result_sheet.iter_rows(values_only=True))
        headers = [str(value or "") for value in rows[0]] if rows else []
        disease_index = headers.index("病种名称") if "病种名称" in headers else None
        disease_values: List[str] = []
        if disease_index is not None:
            for row in rows[1:]:
                value = str(row[disease_index] or "").strip()
                if value:
                    disease_values.append(value)
        disease_text = "\n".join(disease_values)
        known_bad_hits = [term for term in KNOWN_BAD_DISEASE_TERMS if term in disease_text]
        generic_hits = [value for value in disease_values if value in GENERIC_DISEASE_TERMS]
        return {
            "path": str(path),
            "sheet_names": sheet_names,
            "sheet_count": len(sheet_names),
            "result_rows": max(len(rows) - 1, 0),
            "fixed_26_headers": headers == EXCEL_HEADERS,
            "known_bad_disease_hits": known_bad_hits,
            "generic_disease_hits": generic_hits,
            "non_empty_disease_count": len(disease_values),
        }
    finally:
        workbook.close()


def write_json(path: Path, payload: Dict[str, Any]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return path


def selected_ids_from_file(path: Path, count: int) -> str:
    ids = [item.strip() for item in path.read_text(encoding="utf-8").split(",") if item.strip()]
    return ",".join(ids[:count])


def python_executable() -> str:
    return sys.executable
