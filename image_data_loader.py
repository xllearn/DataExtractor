import csv
import json
from html import escape
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import load_workbook

from utils import EXCEL_HEADERS, ensure_dir


RECORD_FIELDS = ["Title", "Content", "AuditTime", "areaname", "SourceURL", "insurancetypename", "info_id"]


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _default_record(batch_id: str, image_path: Path) -> Dict[str, Any]:
    return {
        "Title": f"图片导入测试-{batch_id}",
        "Content": "",
        "AuditTime": "",
        "areaname": "",
        "SourceURL": f"image-import://{batch_id}/{image_path.stem}",
        "insurancetypename": "",
        "info_id": f"{batch_id}-{image_path.stem}",
    }


def _normalize_record(row: Dict[str, Any], batch_id: str, image_path: Path) -> Dict[str, Any]:
    record = _default_record(batch_id, image_path)
    for field in RECORD_FIELDS:
        if _clean(row.get(field)):
            record[field] = _clean(row.get(field))
    if not record["Content"]:
        record["Content"] = "<p></p>"
    record["_batch_id"] = batch_id
    return record


def _records_from_json(path: Path, batch_id: str, image_path: Path) -> List[Dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    items = payload.get("records") if isinstance(payload, dict) else payload
    if not isinstance(items, list):
        raise ValueError("JSON transcript 必须是数组或包含 records 数组")
    return [_normalize_record(item, batch_id, image_path) for item in items if isinstance(item, dict)]


def _records_from_csv(path: Path, batch_id: str, image_path: Path) -> List[Dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [_normalize_record(row, batch_id, image_path) for row in csv.DictReader(handle)]


def _html_table_from_rows(headers: List[str], rows: List[Dict[str, Any]]) -> str:
    table = ["<table>", "<tr>" + "".join(f"<th>{escape(header)}</th>" for header in headers) + "</tr>"]
    for row in rows:
        table.append("<tr>" + "".join(f"<td>{escape(_clean(row.get(header)))}</td>" for header in headers) + "</tr>")
    table.append("</table>")
    return "\n".join(table)


def _records_from_xlsx(path: Path, batch_id: str, image_path: Path) -> List[Dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return []
    headers = [_clean(value) for value in rows[0]]
    data_rows = [
        {headers[index]: raw_row[index] if index < len(raw_row) else "" for index in range(len(headers))}
        for raw_row in rows[1:]
        if any(_clean(value) for value in raw_row)
    ]
    first = data_rows[0] if data_rows else {}
    table_headers = [header for header in EXCEL_HEADERS if header in headers]
    if not table_headers:
        table_headers = headers
    record = _default_record(batch_id, image_path)
    record.update(
        {
            "AuditTime": _clean(first.get("文章时间") or first.get("AuditTime")),
            "areaname": _clean(first.get("地区名称") or first.get("areaname")),
            "SourceURL": f"image-import://{batch_id}/{_clean(first.get('info_id')) or image_path.stem}",
            "insurancetypename": _clean(first.get("保险类型") or first.get("insurancetypename")),
            "info_id": _clean(first.get("info_id")) or f"{batch_id}-{image_path.stem}",
            "Content": _html_table_from_rows(table_headers, data_rows),
            "_batch_id": batch_id,
        }
    )
    return [record]


def _records_from_ocr(image_path: Path, batch_id: str) -> List[Dict[str, Any]]:
    from image_ocr import recognize_image

    text = recognize_image(image_path)
    record = _default_record(batch_id, image_path)
    record["Content"] = "<p>" + escape(text).replace("\n", "<br/>") + "</p>"
    record["_batch_id"] = batch_id
    return [record]


def load_image_records(
    image_path: str | Path,
    transcript_path: str | Path | None = None,
    output_json: str | Path | None = None,
    batch_id: str = "",
) -> List[Dict[str, Any]]:
    image = Path(image_path)
    if not image.exists():
        raise FileNotFoundError(f"图片不存在: {image}")
    batch = batch_id or f"image_import_{__import__('datetime').datetime.now().strftime('%Y%m%d_%H%M%S')}"
    transcript = Path(transcript_path) if transcript_path else None
    if transcript:
        suffix = transcript.suffix.lower()
        if suffix == ".json":
            records = _records_from_json(transcript, batch, image)
        elif suffix == ".csv":
            records = _records_from_csv(transcript, batch, image)
        elif suffix == ".xlsx":
            records = _records_from_xlsx(transcript, batch, image)
        else:
            raise ValueError("transcript_path 仅支持 .json/.csv/.xlsx")
    else:
        records = _records_from_ocr(image, batch)

    if output_json:
        output = Path(output_json)
        ensure_dir(output.parent)
        output.write_text(
            json.dumps({"batch_id": batch, "image_path": str(image), "record_count": len(records), "records": records}, ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )
    return records
