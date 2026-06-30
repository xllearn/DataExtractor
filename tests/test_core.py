import tempfile
import unittest
from pathlib import Path
import sys


sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


class JsonUtilsTests(unittest.TestCase):
    def test_parse_markdown_json_and_apply_defaults(self):
        from json_utils import normalize_llm_rows
        from utils import EXCEL_HEADERS

        record = {
            "Title": "测试标题",
            "SourceURL": "https://example.com/a",
            "AuditTime": "2024-01-02 03:04:05",
            "province": "山东省",
            "areaname": "德州市",
            "insurancetypename": "商业补充保险",
        }
        raw = '```json\n[{"类型":"住院待遇","报销比例":"80%"}]\n```'

        rows = normalize_llm_rows(raw, record, "20260602", Path(tempfile.gettempdir()))

        self.assertEqual(len(rows), 1)
        self.assertEqual(list(rows[0].keys()), EXCEL_HEADERS)
        self.assertEqual(rows[0]["文章时间"], "2024/1/2")
        self.assertEqual(rows[0]["审核日期"], "20260602")
        self.assertEqual(rows[0]["info_id"], "")
        self.assertEqual(rows[0]["地区名称"], "山东省-德州市")
        self.assertEqual(rows[0]["保险类型"], "商业补充保险")
        self.assertEqual(rows[0]["备注"], "--")
        self.assertEqual(rows[0]["相关资讯"], "")
        self.assertEqual(rows[0]["审核状态0待审核1已审核"], 0)
        self.assertEqual(rows[0]["是否需要手动修改执行状态(1是0否)"], 0)
        self.assertEqual(rows[0]["类型"], "住院待遇")
        self.assertEqual(rows[0]["报销比例"], "80%")
        self.assertEqual(rows[0]["起付标准"], "--")

    def test_invalid_json_creates_fallback_and_failed_output(self):
        from json_utils import normalize_llm_rows

        with tempfile.TemporaryDirectory() as tmp:
            record = {
                "Title": "兜底标题",
                "SourceURL": "https://example.com/fallback",
                "AuditTime": "2024-02-03",
            }

            rows = normalize_llm_rows("not json", record, "20260602", Path(tmp))

            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["文章时间"], "2024/2/3")
            self.assertEqual(rows[0]["info_id"], "")
            self.assertEqual(rows[0]["备注"], "--")
            self.assertEqual(rows[0]["相关资讯"], "")
            failed_files = list((Path(tmp) / "failed_llm_outputs").glob("*.txt"))
            self.assertEqual(len(failed_files), 1)


class HtmlParserTests(unittest.TestCase):
    def test_parse_html_extracts_text_table_and_image_urls(self):
        from html_parser import parse_html_content

        html = """
        <html><head><style>.x{}</style><script>alert(1)</script></head>
        <body>
          <p><strong>保障责任</strong>：住院待遇</p>
          <table><tr><th>医院</th><th>比例</th></tr><tr><td>三级</td><td>70%</td></tr></table>
          <img data-src="//cdn.example.com/a.png" />
          <img src="/upload/b.jpg" />
          <img data-original="relative/c.jpg" />
        </body></html>
        """

        parsed = parse_html_content(html, image_base_url="https://img.example.com/root")

        self.assertIn("保障责任", parsed.clean_text)
        self.assertIn("| 医院 | 比例 |", parsed.tables_text)
        self.assertEqual(
            parsed.image_urls,
            [
                "https://cdn.example.com/a.png",
                "https://img.example.com/upload/b.jpg",
                "https://img.example.com/root/relative/c.jpg",
            ],
        )


class ExcelAndInputTests(unittest.TestCase):
    def test_write_rows_creates_expected_headers(self):
        from excel_writer import write_rows_to_workbook
        from openpyxl import load_workbook
        from utils import EXCEL_HEADERS, create_fallback_row

        with tempfile.TemporaryDirectory() as tmp:
            record = {"Title": "A", "AuditTime": "2024-01-01", "SourceURL": "u"}
            rows = [create_fallback_row(record, "20260602")]
            out = Path(tmp) / "out.xlsx"

            write_rows_to_workbook(rows, out, template_path=Path(tmp) / "missing.xlsx")

            wb = load_workbook(out)
            ws = wb.active
            headers = [cell.value for cell in ws[1]]
            self.assertEqual(headers, EXCEL_HEADERS)
            self.assertEqual(ws["A2"].value, "2024/1/1")
            self.assertIsNone(ws["C2"].value)
            self.assertEqual(ws["T2"].value, "--")
            self.assertIsNone(ws["U2"].value)
            self.assertEqual(ws.auto_filter.ref, f"A1:Z{ws.max_row}")
            self.assertEqual(ws.freeze_panes, "A2")

    def test_read_input_xlsx_supports_content_and_limit_offset(self):
        from input_xlsx import read_records_from_xlsx
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "input.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(["Title", "Context", "AuditTime"])
            ws.append(["first", "<p>一</p>", "2024-01-01"])
            ws.append(["second", "<p>二</p>", "2024-01-02"])
            wb.save(path)

            rows = read_records_from_xlsx(path, limit=1, offset=1)

            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["Title"], "second")
            self.assertEqual(rows[0]["Content"], "<p>二</p>")


if __name__ == "__main__":
    unittest.main()
