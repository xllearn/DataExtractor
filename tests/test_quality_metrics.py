import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

import sys


sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def _write_workbook(path: Path, rows):
    from utils import EXCEL_HEADERS

    wb = Workbook()
    ws = wb.active
    ws.title = "结果数据"
    ws.append(EXCEL_HEADERS)
    for row in rows:
        ws.append([row.get(header, "") for header in EXCEL_HEADERS])
    wb.save(path)


class QualityMetricsTests(unittest.TestCase):
    def test_compare_workbooks_normalizes_percent_amount_and_aligns_rows(self):
        from excel_compare import compare_excel_files

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            generated = tmp_path / "generated.xlsx"
            manual = tmp_path / "manual.xlsx"
            report_xlsx = tmp_path / "report.xlsx"
            report_json = tmp_path / "report.json"

            _write_workbook(
                generated,
                [
                    {
                        "info_id": "B-2",
                        "地区名称": "陕西省-西安市",
                        "病种名称": "类风湿性关节炎",
                        "报销比例": "80％",
                        "补助限额": "150000元",
                    },
                    {
                        "info_id": "A-1",
                        "地区名称": "陕西省-西安市",
                        "人员类型": "退休人员",
                        "起付标准": "200元",
                    },
                ],
            )
            _write_workbook(
                manual,
                [
                    {
                        "info_id": "A-1",
                        "地区名称": "陕西省-西安市",
                        "人员类型": "退休人员",
                        "起付标准": "200.00元",
                    },
                    {
                        "info_id": "B-2",
                        "地区名称": "陕西省-西安市",
                        "病种名称": "类风湿性关节炎",
                        "报销比例": "80%",
                        "补助限额": "15万元",
                    },
                ],
            )

            result = compare_excel_files(generated, manual, output_xlsx=report_xlsx, output_json=report_json)

            self.assertGreaterEqual(result["overall_similarity"], 0.95)
            self.assertGreaterEqual(result["core_field_similarity"], 0.95)
            self.assertEqual(result["row_count"]["generated"], 2)
            self.assertEqual(result["row_count"]["manual"], 2)
            self.assertEqual(result["field_metrics"]["报销比例"]["accuracy"], 1.0)
            self.assertEqual(result["field_metrics"]["补助限额"]["accuracy"], 1.0)
            self.assertTrue(report_xlsx.exists())
            self.assertTrue(report_json.exists())

            saved = json.loads(report_json.read_text(encoding="utf-8"))
            self.assertIn("field_metrics", saved)
            wb = load_workbook(report_xlsx)
            try:
                self.assertIn("总体指标", wb.sheetnames)
                self.assertIn("字段指标", wb.sheetnames)
                self.assertIn("差异明细", wb.sheetnames)
            finally:
                wb.close()


if __name__ == "__main__":
    unittest.main()
