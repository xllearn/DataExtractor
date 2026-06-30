import tempfile
import unittest
from pathlib import Path
import sys

from openpyxl import load_workbook


sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


class ExcelMultiSheetTests(unittest.TestCase):
    def test_write_extraction_workbook_creates_traceability_sheets(self):
        from excel_writer import write_extraction_workbook
        from field_mapping import load_field_mapping
        from utils import EXCEL_HEADERS

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "out.xlsx"
            write_extraction_workbook(
                result_rows=[{"报销比例": "80%", "多余": "x"}],
                output_path=output,
                template_path=Path(tmp) / "missing.xlsx",
                field_mapping=load_field_mapping(None),
                collection_logs=[{"source_id": "S1", "status": "success"}],
                field_evidence=[
                    {"source": "database_direct", "field": "地区名称", "value": "山东省"},
                    {"source": "table", "field": "报销比例", "value": "80%"},
                    {"source": "text_rule", "field": "起付标准", "value": "500元"},
                    {"source": "llm", "field": "补助限额", "value": "15万元"},
                ],
                conflict_evidence=[{"field": "报销比例", "rule_value": "80%", "llm_value": "70%"}],
                extract_evaluations=[{"record_index": 1, "confidence_score": 80}],
                failed_records=[{"phase": "process_record", "error": "boom"}],
            )

            wb = load_workbook(output)

        self.assertEqual(
            set(["结果数据", "采集日志", "字段证据", "冲突证据", "抽取评估", "失败记录"]),
            set(wb.sheetnames),
        )
        self.assertEqual([cell.value for cell in wb["结果数据"][1]], EXCEL_HEADERS)
        field_sources = [row[0].value for row in wb["字段证据"].iter_rows(min_row=2, min_col=9, max_col=9)]
        self.assertIn("database_direct", field_sources)
        self.assertIn("table", field_sources)
        self.assertIn("text_rule", field_sources)
        self.assertIn("llm", field_sources)
        self.assertEqual(wb["冲突证据"]["E2"].value, "报销比例")
        self.assertEqual(wb["抽取评估"]["F2"].value, 80)
        self.assertEqual(wb["失败记录"]["A2"].value, "process_record")


if __name__ == "__main__":
    unittest.main()
