from pathlib import Path


def test_workbook_writes_review_confidence_and_row_match_sheets(tmp_path):
    from excel_writer import write_extraction_workbook
    from field_mapping import load_field_mapping
    from openpyxl import load_workbook
    from utils import EXCEL_HEADERS

    out = tmp_path / "out.xlsx"
    write_extraction_workbook(
        [{"类型": "门诊", "报销比例": "80%"}],
        out,
        Path(tmp_path) / "missing.xlsx",
        load_field_mapping(None),
        field_confidence=[
            {
                "row_index": 1,
                "field": "报销比例",
                "value": "80%",
                "confidence": 0.95,
                "source": "table",
                "evidence_count": 1,
                "conflict_count": 0,
                "match_level": "strong",
                "needs_review": False,
                "attempt": "initial",
                "review_status": "pending",
                "reason": "表格坐标证据明确",
                "evidence": "表格",
                "evidence_ids": "E1",
            }
        ],
        review_rows=[
            {
                "source_id": "S1",
                "info_id": "I1",
                "title": "标题",
                "source_url": "https://example.test",
                "row_index": 1,
                "field": "类型",
                "current_value": "",
                "confidence": 0.4,
                "reason": "低置信",
                "evidence": "",
                "evidence_id": "",
                "attempt": "initial",
                "review_status": "pending",
                "reviewed_value": "",
                "review_comment": "",
                "suggested_action": "补充或确认字段",
            }
        ],
        row_match_evidence=[{"source_a": "table", "source_b": "llm", "row_a": 1, "row_b": 2, "similarity": 0.8}],
        candidate_rows=[
            {
                "报销比例": "60%",
                "row_quality_level": "medium",
                "row_quality_reason": "unknown table kept for manual candidate review",
                "mapped_values_json": '{"报销比例": "60%"}',
            }
        ],
        low_value_rows=[
            {
                "source_id": "S1",
                "table_type": "contact_table",
                "row_text": "客服电话 95500",
                "mapped_values_json": '{"报销比例": "70%"}',
                "reason": "obvious noise table",
            }
        ],
        result_index_rows=[{"result_row_index": 1, "table_type": "treatment_table", "target_sheet": "main"}],
        table_classification_rows=[{"table_index": 1, "table_type": "contact_table", "confidence": 0.9}],
    )

    wb = load_workbook(out)
    assert "结果数据" in wb.sheetnames
    assert "字段置信度" in wb.sheetnames
    assert "人工复核" in wb.sheetnames
    assert "行匹配证据" in wb.sheetnames
    assert "结果索引" in wb.sheetnames
    assert "候选结果" in wb.sheetnames
    assert "低价值表格行" in wb.sheetnames
    assert "表格分类" in wb.sheetnames
    assert [cell.value for cell in wb["结果数据"][1]] == EXCEL_HEADERS
    assert len([cell.value for cell in wb["结果数据"][1]]) == 26
    assert len([cell.value for cell in wb["结果数据"][1]]) == wb["结果数据"].max_column
    assert "mapped_values_json" in [cell.value for cell in wb["低价值表格行"][1]]
    assert wb["低价值表格行"].cell(row=2, column=12).value
    assert wb["候选结果"].cell(row=2, column=19).value == "60%"
    assert [cell.value for cell in wb["字段置信度"][1]] == [
        "row_index",
        "field",
        "value",
        "confidence",
        "source",
        "evidence_count",
        "conflict_count",
        "match_level",
        "needs_review",
        "attempt",
        "review_status",
        "reason",
        "evidence",
        "evidence_ids",
    ]
    assert [cell.value for cell in wb["人工复核"][1]] == [
        "source_id",
        "info_id",
        "title",
        "source_url",
        "row_index",
        "field",
        "current_value",
        "confidence",
        "reason",
        "evidence",
        "evidence_id",
        "attempt",
        "review_status",
        "reviewed_value",
        "review_comment",
        "suggested_action",
    ]
    assert wb["人工复核"].cell(row=2, column=13).value == "pending"
