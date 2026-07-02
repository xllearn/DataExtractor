from io import BytesIO
from datetime import datetime
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient

import sys


sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


class ApiReviewTests(unittest.TestCase):
    def _client(self, output_dir: Path) -> TestClient:
        from api_server import create_app
        from excel_writer import write_extraction_workbook
        from field_mapping import load_field_mapping

        records = [
            {
                "_source_id": "S1",
                "info_id": "INFO-1",
                "Title": "复核测试",
                "SourceURL": "https://example.test/review",
                "Content": "<p>报销比例 80%</p>",
            }
        ]

        def provider(keyword="", selected_ids=None, limit=50, offset=0):
            return {"items": records[offset : offset + limit], "total": len(records)}

        def runner(selected_ids, mode, no_ocr, no_llm, external_ocr_text="", prompt_version="v3"):
            output = output_dir / "review_source.xlsx"
            write_extraction_workbook(
                [
                    {
                        "info_id": "INFO-1",
                        "地区名称": "北京",
                        "类型": "门诊",
                        "报销比例": "80%",
                        "备注": "原始备注",
                    }
                ],
                output,
                template_path=output_dir / "missing.xlsx",
                field_mapping=load_field_mapping(None),
                review_rows=[
                    {
                        "source_id": "S1",
                        "info_id": "INFO-1",
                        "title": "复核测试",
                        "source_url": "https://example.test/review?token=secret",
                        "row_index": 1,
                        "field": "报销比例",
                        "current_value": "80%",
                        "confidence": 0.42,
                        "reason": f"低置信度 {output_dir}\\secret-db.txt token=secret",
                        "evidence": f"表格证据 {output_dir}\\evidence.txt",
                        "evidence_id": "E1",
                        "attempt": "initial",
                        "review_status": "pending",
                        "reviewed_value": "",
                        "review_comment": "",
                        "suggested_action": "确认字段",
                    },
                    {
                        "source_id": "S1",
                        "info_id": "INFO-1",
                        "title": "复核测试",
                        "source_url": "https://example.test/review",
                        "row_index": 1,
                        "field": "类型",
                        "current_value": "门诊",
                        "confidence": 0.65,
                        "reason": "来源冲突",
                        "evidence": "规则/LLM",
                        "evidence_id": "E2",
                        "attempt": "initial",
                        "review_status": "pending",
                        "reviewed_value": "",
                        "review_comment": "",
                        "suggested_action": "接受或修改",
                    },
                    {
                        "source_id": "S1",
                        "info_id": "INFO-1",
                        "title": "复核测试",
                        "source_url": "https://example.test/review",
                        "row_index": 1,
                        "field": "备注",
                        "current_value": "原始备注",
                        "confidence": 0.5,
                        "reason": r"噪声字段 C:\Users\admin\My Project\secret-db.txt token=secret",
                        "evidence": r"C:\Users\admin\My Project\evidence file.txt",
                        "evidence_id": "E3",
                        "attempt": "initial",
                        "review_status": "pending",
                        "reviewed_value": "",
                        "review_comment": "",
                        "suggested_action": "可忽略",
                    },
                ],
            )
            return output

        app = create_app(record_provider=provider, extract_runner=runner, output_dir=output_dir, log_dir=output_dir / "logs")
        return TestClient(app)

    def _create_success_job(self, client: TestClient) -> dict:
        response = client.post("/api/extract", json={"selected_ids": ["S1"], "mode": "merge", "no_ocr": True})
        self.assertEqual(response.status_code, 200)
        job_id = response.json()["job_id"]
        payload = {}
        for _ in range(30):
            status = client.get(f"/api/jobs/{job_id}")
            self.assertEqual(status.status_code, 200)
            payload = status.json()
            if payload["status"] in {"success", "failed"}:
                break
            time.sleep(0.05)
        self.assertEqual(payload["status"], "success")
        return payload

    def test_review_items_can_be_updated_applied_and_downloaded(self):
        with tempfile.TemporaryDirectory() as tmp:
            output_dir = Path(tmp)
            client = self._client(output_dir)
            job = self._create_success_job(client)
            job_id = job["job_id"]

            response = client.get(f"/api/jobs/{job_id}/review-items")
            self.assertEqual(response.status_code, 200)
            payload = response.json()
            self.assertEqual(payload["job_id"], job_id)
            self.assertEqual(payload["total"], 3)
            self.assertNotIn(str(output_dir), str(payload))
            self.assertNotIn("secret", str(payload).lower())
            self.assertNotIn("Project", str(payload))
            self.assertNotIn("secret-db", str(payload))
            self.assertNotIn("evidence file", str(payload))
            self.assertIn("https://example.test/review", str(payload))

            again = client.get(f"/api/jobs/{job_id}/review-items").json()
            self.assertEqual([item["item_id"] for item in payload["items"]], [item["item_id"] for item in again["items"]])

            by_field = {item["field"]: item for item in payload["items"]}
            required = {
                "item_id",
                "row_index",
                "field",
                "current_value",
                "reviewed_value",
                "confidence",
                "reason",
                "evidence",
                "status",
                "comment",
                "updated_at",
            }
            self.assertTrue(required.issubset(by_field["报销比例"].keys()))
            self.assertEqual(by_field["报销比例"]["status"], "pending")

            edit = client.post(
                f"/api/jobs/{job_id}/review-items/{by_field['报销比例']['item_id']}",
                json={"status": "edited", "reviewed_value": "90%", "comment": "人工确认 90%"},
            )
            self.assertEqual(edit.status_code, 200)
            self.assertEqual(edit.json()["status"], "edited")
            self.assertEqual(edit.json()["reviewed_value"], "90%")

            accept = client.post(
                f"/api/jobs/{job_id}/review-items/{by_field['类型']['item_id']}",
                json={"status": "accepted", "comment": "原值可用"},
            )
            self.assertEqual(accept.status_code, 200)
            self.assertEqual(accept.json()["reviewed_value"], "门诊")

            ignored = client.post(
                f"/api/jobs/{job_id}/review-items/{by_field['备注']['item_id']}",
                json={"status": "ignored", "comment": "本轮不处理"},
            )
            self.assertEqual(ignored.status_code, 200)
            self.assertEqual(ignored.json()["status"], "ignored")

            apply_response = client.post(f"/api/jobs/{job_id}/apply-reviews")
            self.assertEqual(apply_response.status_code, 200)
            applied = apply_response.json()
            self.assertEqual(applied["job_id"], job_id)
            self.assertEqual(applied["applied_count"], 3)
            self.assertEqual(applied["edited_count"], 1)
            self.assertTrue(applied["reviewed_workbook_url"].endswith("/reviewed-workbook"))
            self.assertNotIn("path", applied)
            self.assertNotIn(str(output_dir), str(applied))

            download = client.get(applied["reviewed_workbook_url"])
            self.assertEqual(download.status_code, 200)
            self.assertIn("spreadsheetml", download.headers["content-type"])
            self.assertIn("_reviewed", download.headers.get("content-disposition", ""))

            from openpyxl import load_workbook

            workbook = load_workbook(BytesIO(download.content))
            self.assertIn("结果数据", workbook.sheetnames)
            self.assertIn("人工复核", workbook.sheetnames)
            self.assertIn("复核日志", workbook.sheetnames)

            result_sheet = workbook["结果数据"]
            result_headers = [cell.value for cell in result_sheet[1]]
            reimbursement_column = result_headers.index("报销比例") + 1
            self.assertEqual(result_sheet.cell(row=2, column=reimbursement_column).value, "90%")

            review_sheet = workbook["人工复核"]
            review_headers = [cell.value for cell in review_sheet[1]]
            status_column = review_headers.index("review_status") + 1
            value_column = review_headers.index("reviewed_value") + 1
            comment_column = review_headers.index("review_comment") + 1
            field_column = review_headers.index("field") + 1
            review_rows = {
                review_sheet.cell(row=row_index, column=field_column).value: row_index
                for row_index in range(2, review_sheet.max_row + 1)
            }
            edited_row = review_rows["报销比例"]
            self.assertEqual(review_sheet.cell(row=edited_row, column=status_column).value, "edited")
            self.assertEqual(review_sheet.cell(row=edited_row, column=value_column).value, "90%")
            self.assertEqual(review_sheet.cell(row=edited_row, column=comment_column).value, "人工确认 90%")
            self.assertEqual(review_sheet.cell(row=review_rows["类型"], column=status_column).value, "accepted")
            self.assertEqual(review_sheet.cell(row=review_rows["备注"], column=status_column).value, "ignored")

            log_sheet = workbook["复核日志"]
            log_headers = [cell.value for cell in log_sheet[1]]
            self.assertIn("item_id", log_headers)
            self.assertIn("status", log_headers)
            self.assertEqual(log_sheet.max_row, 4)
            workbook.close()

            with patch("services.review_service.datetime") as fake_datetime:
                fake_datetime.now.return_value = datetime(2026, 7, 2, 12, 0, 0)
                first = client.post(f"/api/jobs/{job_id}/apply-reviews").json()
                second = client.post(f"/api/jobs/{job_id}/apply-reviews").json()
            self.assertNotEqual(first["reviewed_file_name"], second["reviewed_file_name"])

    def test_review_errors_are_json_and_bound_to_successful_jobs(self):
        with tempfile.TemporaryDirectory() as tmp:
            output_dir = Path(tmp)
            client = self._client(output_dir)
            job = self._create_success_job(client)
            job_id = job["job_id"]

            missing_workbook = client.get(f"/api/jobs/{job_id}/reviewed-workbook")
            self.assertEqual(missing_workbook.status_code, 404)
            self.assertEqual(missing_workbook.headers["content-type"].split(";")[0], "application/json")

            unknown_item = client.post(f"/api/jobs/{job_id}/review-items/rev_missing", json={"status": "accepted"})
            self.assertEqual(unknown_item.status_code, 404)
            self.assertEqual(unknown_item.headers["content-type"].split(";")[0], "application/json")

            items_response = client.get(f"/api/jobs/{job_id}/review-items")
            self.assertEqual(items_response.status_code, 200)
            item_id = items_response.json()["items"][0]["item_id"]
            invalid_status = client.post(f"/api/jobs/{job_id}/review-items/{item_id}", json={"status": "done"})
            self.assertEqual(invalid_status.status_code, 400)
            self.assertEqual(invalid_status.headers["content-type"].split(";")[0], "application/json")

            from services.job_store import JobStore

            running_job_id = "job_20260702_120000_abcdef12"
            JobStore(output_dir).create(job_id=running_job_id, status="running", file_id=job["file_id"])
            not_ready = client.get(f"/api/jobs/{running_job_id}/reviewed-workbook")
            self.assertEqual(not_ready.status_code, 400)
            self.assertEqual(not_ready.headers["content-type"].split(";")[0], "application/json")
