from datetime import datetime
from pathlib import Path
from copy import copy
from typing import Any, Dict, Iterable, List

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from field_mapping import FieldMapping, load_field_mapping, normalize_record_fields
from security_utils import sanitize_excel_value
from utils import EXCEL_HEADERS, build_single_filename, ensure_dir


LONG_TEXT_HEADERS = {"个人账户计入办法", "个人账户使用范围", "备注", "相关资讯"}
RESULT_INDEX_SHEET = "结果索引"
CANDIDATE_RESULT_SHEET = "候选结果"
LOW_VALUE_TABLE_ROW_SHEET = "低价值表格行"
TABLE_CLASSIFICATION_SHEET = "表格分类"
TARGET_SHEETS = {
    "结果数据",
    "采集日志",
    "字段证据",
    "冲突证据",
    "抽取评估",
    "失败记录",
    "字段置信度",
    "人工复核",
    "行匹配证据",
    RESULT_INDEX_SHEET,
    CANDIDATE_RESULT_SHEET,
    LOW_VALUE_TABLE_ROW_SHEET,
    TABLE_CLASSIFICATION_SHEET,
}


def _new_or_template_workbook(template_path: Path) -> Workbook:
    if template_path.exists():
        workbook = load_workbook(template_path)
    else:
        workbook = Workbook()
    for worksheet in list(workbook.worksheets):
        if worksheet.title in TARGET_SHEETS or (len(workbook.worksheets) == 1 and worksheet.title == "Sheet"):
            workbook.remove(worksheet)
    return workbook


def write_rows_to_workbook(
    rows: Iterable[Dict[str, object]],
    output_path: Path,
    template_path: Path,
    field_mapping: FieldMapping | None = None,
) -> Path:
    return write_extraction_workbook(
        result_rows=list(rows),
        output_path=output_path,
        template_path=template_path,
        field_mapping=field_mapping or load_field_mapping(None),
    )


def write_extraction_workbook(
    result_rows: List[Dict[str, Any]],
    output_path: Path,
    template_path: Path,
    field_mapping: FieldMapping,
    collection_logs: List[Dict[str, Any]] | None = None,
    field_evidence: List[Dict[str, Any]] | None = None,
    conflict_evidence: List[Dict[str, Any]] | None = None,
    extract_evaluations: List[Dict[str, Any]] | None = None,
    failed_records: List[Dict[str, Any]] | None = None,
    field_confidence: List[Dict[str, Any]] | None = None,
    review_rows: List[Dict[str, Any]] | None = None,
    row_match_evidence: List[Dict[str, Any]] | None = None,
    candidate_rows: List[Dict[str, Any]] | None = None,
    low_value_rows: List[Dict[str, Any]] | None = None,
    result_index_rows: List[Dict[str, Any]] | None = None,
    table_classification_rows: List[Dict[str, Any]] | None = None,
) -> Path:
    output_path = Path(output_path)
    ensure_dir(output_path.parent)
    workbook = _new_or_template_workbook(template_path)

    result_sheet = workbook.create_sheet("结果数据")
    ensure_headers(result_sheet, field_mapping)
    append_rows(result_sheet, result_rows, field_mapping)
    format_worksheet(result_sheet, field_mapping)

    write_dict_sheet(workbook, "采集日志", collection_logs or [], COLLECTION_LOG_HEADERS)
    write_dict_sheet(workbook, "字段证据", field_evidence or [], FIELD_EVIDENCE_HEADERS)
    write_dict_sheet(workbook, "冲突证据", conflict_evidence or [], CONFLICT_EVIDENCE_HEADERS)
    write_dict_sheet(workbook, "抽取评估", extract_evaluations or [], EXTRACT_EVAL_HEADERS)
    if field_confidence is not None:
        write_dict_sheet(workbook, "字段置信度", field_confidence, FIELD_CONFIDENCE_HEADERS)
    if review_rows is not None:
        write_dict_sheet(workbook, "人工复核", review_rows, REVIEW_ROW_HEADERS)
    if row_match_evidence is not None:
        write_dict_sheet(workbook, "行匹配证据", row_match_evidence, ROW_MATCH_HEADERS)
    write_dict_sheet(workbook, RESULT_INDEX_SHEET, result_index_rows or [], RESULT_INDEX_HEADERS)
    write_dict_sheet(workbook, CANDIDATE_RESULT_SHEET, candidate_rows or [], [*field_mapping.headers, *CANDIDATE_EXTRA_HEADERS])
    write_dict_sheet(workbook, LOW_VALUE_TABLE_ROW_SHEET, low_value_rows or [], LOW_VALUE_HEADERS)
    write_dict_sheet(workbook, TABLE_CLASSIFICATION_SHEET, table_classification_rows or [], TABLE_CLASSIFICATION_HEADERS)
    write_dict_sheet(workbook, "失败记录", failed_records or [], FAILED_RECORD_HEADERS)

    workbook.save(output_path)
    return output_path


def ensure_headers(worksheet, field_mapping: FieldMapping | None = None) -> None:
    headers = (field_mapping or load_field_mapping(None)).headers
    for index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=index)
        cell.value = header
        if cell.font:
            font = copy(cell.font)
            font.bold = True
            cell.font = font
        else:
            cell.font = Font(bold=True)
        if cell.fill is None or cell.fill.fill_type is None:
            cell.fill = PatternFill("solid", fgColor="D9EAF7")


def append_rows(worksheet, rows: Iterable[Dict[str, object]], field_mapping: FieldMapping | None = None) -> None:
    field_mapping = field_mapping or load_field_mapping(None)
    for row in rows:
        normalized = normalize_record_fields(row, field_mapping)
        worksheet.append([sanitize_excel_value(normalized.get(header, "")) for header in field_mapping.headers])


def format_worksheet(worksheet, field_mapping: FieldMapping | None = None) -> None:
    headers = (field_mapping or load_field_mapping(None)).headers
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{worksheet.max_row}"

    for column_index, header in enumerate(headers, start=1):
        letter = get_column_letter(column_index)
        max_length = len(header)
        for cell in worksheet[letter]:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, min(len(value), 60))
            if header in LONG_TEXT_HEADERS or len(value) > 30:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                cell.alignment = Alignment(vertical="top")
        worksheet.column_dimensions[letter].width = min(max(max_length + 2, 10), 45)


COLLECTION_LOG_HEADERS = [
    "source_id",
    "info_id",
    "title",
    "source_url",
    "status",
    "attempt",
    "input_mode",
    "ocr_available",
    "ocr_status_reason",
    "vision_enabled",
    "vision_triggered",
    "vision_success_count",
    "vision_failure_count",
    "vision_model",
    "vision_error",
    "ocr_triggered",
    "ocr_trigger_reason",
    "ocr_skipped_reason",
    "image_count",
    "ocr_success_count",
    "ocr_failure_count",
    "ocr_failure_reason",
    "external_ocr_used",
    "llm_format",
    "prompt_version",
    "prompt_hash",
    "response_hash",
    "llm_parse_success",
    "need_manual_review",
    "review_reason",
    "output_rows",
    "error",
    "initial_confidence_score",
    "ocr_retry_confidence_score",
    "ocr_improved",
    "final_attempt",
]
FIELD_EVIDENCE_HEADERS = [
    "source_id",
    "info_id",
    "record_index",
    "row_index",
    "field",
    "value",
    "evidence",
    "confidence",
    "source",
    "rule_name",
    "attempt",
    "chosen",
    "table_index",
    "col_index",
    "header",
    "header_path",
    "cell_text",
    "caption",
    "evidence_id",
]
CONFLICT_EVIDENCE_HEADERS = [
    "source_id",
    "info_id",
    "record_index",
    "row_index",
    "field",
    "rule_value",
    "llm_value",
    "chosen_value",
    "reason",
    "rule_source",
    "llm_source",
    "source_a",
    "value_a",
    "source_b",
    "value_b",
    "chosen_source",
    "attempt",
]
EXTRACT_EVAL_HEADERS = [
    "record_index",
    "row_index",
    "attempt",
    "Title",
    "SourceURL",
    "confidence_score",
    "confidence_level",
    "confidence_reason",
    "should_retry_with_ocr",
    "ocr_trigger_reason",
    "evidence_score",
    "key_field_score",
    "ocr_risk_score",
]
FIELD_CONFIDENCE_HEADERS = [
    "row_index",
    "field",
    "value",
    "confidence",
    "source",
    "evidence_count",
    "conflict_count",
    "match_level",
    "needs_review",
    "attempt",
    "review_status",
    "reason",
    "evidence",
    "evidence_ids",
]
REVIEW_ROW_HEADERS = [
    "source_id",
    "info_id",
    "title",
    "source_url",
    "row_index",
    "field",
    "current_value",
    "confidence",
    "reason",
    "evidence",
    "evidence_id",
    "attempt",
    "review_status",
    "reviewed_value",
    "review_comment",
    "suggested_action",
]
ROW_MATCH_HEADERS = ["record_index", "attempt", "source_a", "source_b", "row_a", "row_b", "similarity", "matched_fields", "reason", "candidate_index"]
FAILED_RECORD_HEADERS = ["phase", "record_index", "source_id", "info_id", "Title", "SourceURL", "error"]
RESULT_INDEX_HEADERS = [
    "result_row_index",
    "source_id",
    "info_id",
    "title",
    "source_url",
    "extraction_source",
    "table_index",
    "row_index",
    "table_type",
    "table_confidence",
    "table_confidence_level",
    "row_quality_level",
    "row_quality_reason",
    "has_treatment_context_source",
    "target_sheet",
]
CANDIDATE_EXTRA_HEADERS = [
    "row_quality_level",
    "row_quality_reason",
    "has_treatment_context",
    "has_treatment_context_source",
    "table_type",
    "table_confidence",
    "table_confidence_level",
    "source_id",
    "info_id",
    "title",
    "source_url",
    "mapped_values_json",
]
LOW_VALUE_HEADERS = [
    "source_id",
    "info_id",
    "title",
    "source_url",
    "table_index",
    "row_index",
    "table_type",
    "table_confidence",
    "table_confidence_level",
    "headers",
    "row_text",
    "mapped_values_json",
    "reason",
]
TABLE_CLASSIFICATION_HEADERS = [
    "source_id",
    "info_id",
    "title",
    "table_index",
    "table_type",
    "confidence",
    "confidence_level",
    "positive_signals",
    "negative_signals",
    "reason",
    "headers",
    "caption",
    "row_count",
]


def write_dict_sheet(workbook, title: str, rows: List[Dict[str, Any]], headers: List[str]) -> None:
    worksheet = workbook.create_sheet(title)
    for index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=index)
        cell.value = header
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows:
        worksheet.append([sanitize_excel_value(row.get(header, "")) for header in headers])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{worksheet.max_row}"
    for column_index, header in enumerate(headers, start=1):
        letter = get_column_letter(column_index)
        max_length = len(header)
        for cell in worksheet[letter]:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, min(len(value), 60))
            cell.alignment = Alignment(wrap_text=len(value) > 30, vertical="top")
        worksheet.column_dimensions[letter].width = min(max(max_length + 2, 10), 45)


def build_single_output_path(record: Dict[str, object], output_dir: Path, sequence: int) -> Path:
    return Path(output_dir) / build_single_filename(record, sequence)


def build_merge_output_path(output_dir: Path) -> Path:
    filename = f"商业补充保险抽取结果_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return Path(output_dir) / filename
