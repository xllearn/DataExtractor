from pathlib import Path
from typing import Any, Dict, List

from openpyxl import load_workbook

from utils import DB_FIELDS


def _normalize_header(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _row_is_empty(values: List[Any]) -> bool:
    return all(value is None or str(value).strip() == "" for value in values)


def read_records_from_xlsx(path: Path, limit: int = 1, offset: int = 0) -> List[Dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    records: List[Dict[str, Any]] = []

    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            try:
                header_row = next(rows)
            except StopIteration:
                continue
            headers = [_normalize_header(value) for value in header_row]
            if not any(headers):
                continue

            data_index = 0
            for row_values in rows:
                values = list(row_values)
                if _row_is_empty(values):
                    continue
                if data_index < offset:
                    data_index += 1
                    continue
                if len(records) >= limit:
                    return records

                record = {header: values[index] if index < len(values) else None for index, header in enumerate(headers) if header}
                if "Content" not in record and "Context" in record:
                    record["Content"] = record.get("Context")
                for field in DB_FIELDS:
                    record.setdefault(field, None)
                records.append(record)
                data_index += 1

            if records:
                return records
        return records
    finally:
        workbook.close()
