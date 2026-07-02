import hashlib
import json
import re
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook

from config_loader import DbConfig, load_db_config
from db_reader import DbReaderError, quote_identifier, quote_table, validate_column_name, validate_table_name
from security_utils import mask_sensitive_text
from services.job_store import JOB_ID_RE
from utils import ensure_dir


RESULT_SHEET = "结果数据"
PREVIEW_ID_RE = re.compile(r"^writeback_\d{8}_\d{6}_[0-9a-f]{8}$")
WINDOWS_PATH_RE = re.compile(r"[A-Za-z]:\\.*?(?=\s+(?:and|or)\s+|[,;\"'\]}]|$)", re.IGNORECASE)
POSIX_PATH_RE = re.compile(r"(?<![\w:/])/(?:[^/\s,;\"'\]}]+/)+[^/\s,;\"'\]}]+")


class WritebackServiceError(ValueError):
    def __init__(self, status_code: int, detail: str, error_type: str):
        super().__init__(detail)
        self.status_code = status_code
        self.detail = detail
        self.error_type = error_type


@dataclass(frozen=True)
class ActiveWritebackConfig:
    database_url: str
    target_table: str
    key_column: str
    allowed_columns: list[str]
    fingerprint: str


class WritebackService:
    def __init__(
        self,
        output_root: str | Path,
        config_path: str | Path = "config/db_config.yml",
        *,
        config_loader: Callable[[str | Path], DbConfig] = load_db_config,
        engine_factory: Callable[[str], Any] | None = None,
    ):
        self.output_root = Path(output_root).resolve()
        self.root = (self.output_root / "writeback").resolve()
        self.config_path = config_path
        self.config_loader = config_loader
        self.engine_factory = engine_factory

    def preview(self, job_id: str, workbook_path: str | Path, dry_run: bool = True) -> dict:
        try:
            dry_run = _strict_bool(dry_run, "dry_run")
            workbook = self._safe_workbook_path(workbook_path)
            config = self._active_config()
            read_result = self._read_workbook_rows(workbook, config)
            row_count = self._count_target_rows(config, read_result["keys"])
            preview_id = _new_preview_id()
            state = {
                "preview_id": preview_id,
                "job_id": job_id,
                "config_fingerprint": config.fingerprint,
                "target_table": config.target_table,
                "key_column": config.key_column,
                "allowed_columns": config.allowed_columns,
                "update_columns": read_result["update_columns"],
                "skipped_columns": read_result["skipped_columns"],
                "row_count": row_count,
                "workbook_row_count": read_result["workbook_row_count"],
                "dry_run": dry_run,
                "workbook_digest": self._file_digest(workbook),
                "created_at": _now_iso(),
            }
            self._write_preview_state(job_id, preview_id, state)
            self._audit(job_id, "preview", "success", state, required=True)
            return self._public_payload(state)
        except WritebackServiceError as exc:
            if exc.error_type != "WritebackAuditFailed":
                self._audit(job_id, "preview", "error", {"error_type": exc.error_type, "detail": exc.detail}, required=False)
            raise

    def commit(self, job_id: str, workbook_path: str | Path, preview_id: str, dry_run: bool = True) -> dict:
        try:
            dry_run = _strict_bool(dry_run, "dry_run")
            workbook = self._safe_workbook_path(workbook_path)
            state = self._read_preview_state(job_id, preview_id)
            if state.get("job_id") != job_id:
                raise WritebackServiceError(409, "writeback preview does not belong to this job", "WritebackPreviewMismatch")

            config = self._active_config()
            if state.get("config_fingerprint") != config.fingerprint:
                raise WritebackServiceError(409, "writeback configuration changed after preview", "WritebackPreviewMismatch")
            if state.get("workbook_digest") != self._file_digest(workbook):
                raise WritebackServiceError(409, "job workbook changed after preview", "WritebackPreviewMismatch")

            read_result = self._read_workbook_rows(workbook, config)
            expected_update_columns = list(state.get("update_columns") or [])
            if read_result["update_columns"] != expected_update_columns:
                raise WritebackServiceError(409, "writeback columns changed after preview", "WritebackPreviewMismatch")

            if dry_run:
                committed_rows = 0
            else:
                start_payload = dict(state)
                start_payload.update({"dry_run": False, "committed_rows": 0})
                self._audit(job_id, "commit", "started", start_payload, required=True)
                committed_rows = self._update_rows(config, read_result["rows"], expected_update_columns)
            payload = dict(state)
            payload.update(
                {
                    "dry_run": dry_run,
                    "committed_rows": committed_rows,
                    "committed_at": _now_iso(),
                }
            )
            self._audit(job_id, "commit", "success", payload, required=True)
            return self._public_payload(payload)
        except WritebackServiceError as exc:
            if exc.error_type != "WritebackAuditFailed":
                self._audit(
                    job_id,
                    "commit",
                    "error",
                    {"preview_id": preview_id, "error_type": exc.error_type, "detail": exc.detail},
                    required=False,
                )
            raise

    def _active_config(self) -> ActiveWritebackConfig:
        try:
            db_config = self.config_loader(self.config_path)
        except Exception as exc:
            raise WritebackServiceError(400, f"writeback config load failed: {self._safe_text(exc)}", "WritebackConfigInvalid") from exc

        writeback = getattr(db_config, "writeback", None)
        if not writeback or not bool(getattr(writeback, "enabled", False)):
            raise WritebackServiceError(400, "writeback is disabled; set writeback.enabled=true to use this API", "WritebackDisabled")
        if not db_config.database_url:
            raise WritebackServiceError(400, "database.url is required for writeback", "WritebackConfigInvalid")
        if not getattr(writeback, "target_table", ""):
            raise WritebackServiceError(400, "writeback.target_table is required", "WritebackConfigInvalid")
        if not getattr(writeback, "key_column", ""):
            raise WritebackServiceError(400, "writeback.key_column is required", "WritebackConfigInvalid")

        try:
            target_table = validate_table_name(str(writeback.target_table))
            key_column = validate_column_name(str(writeback.key_column), "writeback.key_column")
            allowed_columns = _dedupe(
                validate_column_name(str(column), "writeback.allowed_columns")
                for column in list(getattr(writeback, "allowed_columns", []) or [])
            )
        except DbReaderError as exc:
            raise WritebackServiceError(400, self._safe_text(exc), "WritebackConfigInvalid") from exc
        if not allowed_columns:
            raise WritebackServiceError(400, "writeback.allowed_columns must not be empty", "WritebackConfigInvalid")

        fingerprint = self._config_fingerprint(db_config.database_url, target_table, key_column, allowed_columns)
        return ActiveWritebackConfig(
            database_url=db_config.database_url,
            target_table=target_table,
            key_column=key_column,
            allowed_columns=allowed_columns,
            fingerprint=fingerprint,
        )

    def _read_workbook_rows(self, workbook_path: Path, config: ActiveWritebackConfig) -> dict[str, Any]:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        try:
            if RESULT_SHEET not in workbook.sheetnames:
                raise WritebackServiceError(404, "result sheet not found", "WritebackWorkbookInvalid")
            sheet = workbook[RESULT_SHEET]
            headers = self._headers(sheet)
            if config.key_column not in headers:
                raise WritebackServiceError(400, "writeback key column not found in result workbook", "WritebackWorkbookInvalid")
            update_columns = [header for header in headers if header in config.allowed_columns and header != config.key_column]
            skipped_columns = [header for header in headers if header and header != config.key_column and header not in update_columns]
            rows_by_key: dict[str, dict[str, Any]] = {}
            for row in sheet.iter_rows(min_row=2, values_only=True):
                row_data = {headers[index]: row[index] if index < len(row) else None for index in range(len(headers))}
                key_value = "" if row_data.get(config.key_column) is None else str(row_data.get(config.key_column)).strip()
                if not key_value:
                    continue
                rows_by_key[key_value] = {column: row_data.get(column) for column in update_columns}
            rows = [{"key": key, "values": values} for key, values in rows_by_key.items()]
            return {
                "rows": rows,
                "keys": [row["key"] for row in rows],
                "update_columns": update_columns,
                "skipped_columns": skipped_columns,
                "workbook_row_count": len(rows),
            }
        finally:
            workbook.close()

    def _count_target_rows(self, config: ActiveWritebackConfig, keys: list[str]) -> int:
        if not keys:
            return 0
        table = quote_table(config.target_table)
        key_column = quote_identifier(config.key_column)
        params = {f"key_{index}": key for index, key in enumerate(keys)}
        placeholders = ", ".join(f":key_{index}" for index in range(len(keys)))
        statement = f"SELECT COUNT(*) FROM {table} WHERE {key_column} IN ({placeholders})"
        try:
            from sqlalchemy import text

            engine = self._engine(config.database_url)
            try:
                with engine.connect() as connection:
                    return int(connection.execute(text(statement), params).scalar() or 0)
            finally:
                self._dispose_engine(engine)
        except Exception as exc:
            raise WritebackServiceError(400, f"writeback preview failed: {self._safe_text(exc)}", "WritebackPreviewFailed") from exc

    def _update_rows(self, config: ActiveWritebackConfig, rows: list[dict[str, Any]], update_columns: list[str]) -> int:
        if not rows or not update_columns:
            return 0
        table = quote_table(config.target_table)
        key_column = quote_identifier(config.key_column)
        set_expr = ", ".join(f"{quote_identifier(column)} = :value_{index}" for index, column in enumerate(update_columns))
        statement = f"UPDATE {table} SET {set_expr} WHERE {key_column} = :key_value"
        try:
            from sqlalchemy import text

            engine = self._engine(config.database_url)
            try:
                committed_rows = 0
                with engine.begin() as connection:
                    for row in rows:
                        params = {f"value_{index}": row["values"].get(column) for index, column in enumerate(update_columns)}
                        params["key_value"] = row["key"]
                        result = connection.execute(text(statement), params)
                        committed_rows += max(int(result.rowcount or 0), 0)
                return committed_rows
            finally:
                self._dispose_engine(engine)
        except Exception as exc:
            raise WritebackServiceError(400, f"writeback commit failed: {self._safe_text(exc)}", "WritebackCommitFailed") from exc

    def _engine(self, database_url: str):
        if self.engine_factory is not None:
            return self.engine_factory(database_url)
        try:
            from sqlalchemy import create_engine
        except Exception as exc:  # pragma: no cover - dependency guard
            raise WritebackServiceError(500, "SQLAlchemy dependency is missing", "DependencyMissing") from exc
        return create_engine(database_url)

    def _dispose_engine(self, engine: Any) -> None:
        dispose = getattr(engine, "dispose", None)
        if callable(dispose):
            dispose()

    def _safe_workbook_path(self, workbook_path: str | Path) -> Path:
        path = Path(workbook_path).resolve()
        if not self._is_relative_to(path, self.output_root) or path.suffix.lower() != ".xlsx" or not path.exists():
            raise WritebackServiceError(404, "job workbook not found", "JobWorkbookNotFound")
        return path

    def _job_dir(self, job_id: str) -> Path:
        if not JOB_ID_RE.fullmatch(str(job_id or "")):
            raise WritebackServiceError(400, "invalid job_id", "ValidationError")
        path = (self.root / job_id).resolve()
        if not self._is_relative_to(path, self.root):
            raise WritebackServiceError(400, "invalid job_id", "ValidationError")
        return path

    def _preview_path(self, job_id: str, preview_id: str) -> Path:
        if not PREVIEW_ID_RE.fullmatch(str(preview_id or "")):
            raise WritebackServiceError(404, "writeback preview not found", "WritebackPreviewNotFound")
        return self._job_dir(job_id) / f"{preview_id}.json"

    def _read_preview_state(self, job_id: str, preview_id: str) -> dict[str, Any]:
        path = self._preview_path(job_id, preview_id)
        if not path.exists():
            raise WritebackServiceError(404, "writeback preview not found", "WritebackPreviewNotFound")
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except Exception as exc:
            raise WritebackServiceError(500, f"writeback preview read failed: {self._safe_text(exc)}", "WritebackPreviewReadError") from exc
        if not isinstance(payload, dict) or payload.get("preview_id") != preview_id:
            raise WritebackServiceError(500, "writeback preview metadata is invalid", "WritebackPreviewReadError")
        return payload

    def _write_preview_state(self, job_id: str, preview_id: str, state: dict[str, Any]) -> None:
        path = self._preview_path(job_id, preview_id)
        ensure_dir(path.parent)
        payload = self._public_payload(state)
        payload["config_fingerprint"] = str(state.get("config_fingerprint") or "")
        payload["workbook_digest"] = str(state.get("workbook_digest") or "")
        temp_path = path.with_suffix(".json.tmp")
        temp_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
        temp_path.replace(path)

    def _audit(self, job_id: str, event: str, status: str, payload: dict[str, Any], *, required: bool = False) -> None:
        try:
            directory = self._job_dir(job_id)
            ensure_dir(directory)
            row = self._sanitize(
                {
                    "event": event,
                    "status": status,
                    "job_id": job_id,
                    "preview_id": payload.get("preview_id", ""),
                    "row_count": payload.get("row_count", 0),
                    "workbook_row_count": payload.get("workbook_row_count", 0),
                    "committed_rows": payload.get("committed_rows", 0),
                    "dry_run": payload.get("dry_run", True),
                    "error_type": payload.get("error_type", ""),
                    "detail": payload.get("detail", ""),
                    "created_at": _now_iso(),
                }
            )
            with (directory / "audit.jsonl").open("a", encoding="utf-8") as handle:
                handle.write(json.dumps(row, ensure_ascii=False, default=str) + "\n")
        except Exception as exc:
            if required:
                raise WritebackServiceError(500, f"writeback audit failed: {self._safe_text(exc)}", "WritebackAuditFailed") from exc

    def _public_payload(self, payload: dict[str, Any]) -> dict[str, Any]:
        public = {
            "job_id": payload.get("job_id", ""),
            "preview_id": payload.get("preview_id", ""),
            "target_table": payload.get("target_table", ""),
            "key_column": payload.get("key_column", ""),
            "allowed_columns": list(payload.get("allowed_columns") or []),
            "update_columns": list(payload.get("update_columns") or []),
            "skipped_columns": list(payload.get("skipped_columns") or []),
            "row_count": _safe_int(payload.get("row_count")),
            "workbook_row_count": _safe_int(payload.get("workbook_row_count")),
            "committed_rows": _safe_int(payload.get("committed_rows")),
            "dry_run": bool(payload.get("dry_run", True)),
            "created_at": payload.get("created_at", ""),
            "committed_at": payload.get("committed_at", ""),
        }
        return self._sanitize(public)

    def _headers(self, worksheet) -> list[str]:
        if worksheet.max_row < 1:
            return []
        return ["" if cell.value is None else str(cell.value).strip() for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]

    def _config_fingerprint(self, database_url: str, target_table: str, key_column: str, allowed_columns: list[str]) -> str:
        payload = {
            "database_url": database_url,
            "target_table": target_table,
            "key_column": key_column,
            "allowed_columns": allowed_columns,
        }
        return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()

    def _file_digest(self, path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    def _sanitize(self, payload: Any) -> Any:
        if isinstance(payload, dict):
            return {self._safe_text(key): self._sanitize(value) for key, value in payload.items()}
        if isinstance(payload, list):
            return [self._sanitize(item) for item in payload]
        if isinstance(payload, str):
            return self._safe_text(payload)
        return payload

    def _safe_text(self, value: Any) -> str:
        text = mask_sensitive_text("" if value is None else str(value))
        text = WINDOWS_PATH_RE.sub("[path]", text)
        text = POSIX_PATH_RE.sub("[path]", text)
        return text

    def _is_relative_to(self, path: Path, root: Path) -> bool:
        try:
            path.relative_to(root)
            return True
        except ValueError:
            return False


def _new_preview_id() -> str:
    return f"writeback_{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:8]}"


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _safe_int(value: Any) -> int:
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _strict_bool(value: Any, name: str) -> bool:
    if isinstance(value, bool):
        return value
    raise WritebackServiceError(400, f"{name} must be a JSON boolean", "ValidationError")


def _dedupe(values) -> list[str]:
    result: list[str] = []
    for value in values:
        if value and value not in result:
            result.append(value)
    return result
