import hashlib
import json
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from security_utils import mask_sensitive_text, mask_url, sanitize_excel_value
from services.job_store import JOB_ID_RE
from utils import ensure_dir


REVIEW_SHEET = "人工复核"
RESULT_SHEET = "结果数据"
REVIEW_LOG_SHEET = "复核日志"
REVIEW_STATUSES = {"pending", "accepted", "edited", "ignored"}
STATE_FILENAME = "review_state.json"
WINDOWS_PATH_RE = re.compile(r"[A-Za-z]:\\.*?(?=\s+(?:and|or)\s+|[,;\"'\]}]|$)", re.IGNORECASE)
POSIX_PATH_RE = re.compile(r"(?<![\w:])/(?:[^/\s,;\"'\]}]+/)+[^/\s,;\"'\]}]+")


class ReviewServiceError(ValueError):
    def __init__(self, status_code: int, detail: str, error_type: str):
        super().__init__(detail)
        self.status_code = status_code
        self.detail = detail
        self.error_type = error_type


class ReviewService:
    def __init__(self, output_root: str | Path):
        self.output_root = Path(output_root).resolve()
        self.root = (self.output_root / "reviews").resolve()

    def list_items(self, job_id: str, workbook_path: str | Path) -> dict:
        items = self._items_with_state(job_id, workbook_path)
        public_items = [self._public_item(item) for item in items]
        return {"job_id": job_id, "items": public_items, "total": len(public_items)}

    def update_item(
        self,
        job_id: str,
        workbook_path: str | Path,
        item_id: str,
        status: str,
        reviewed_value: Any = None,
        comment: Any = "",
    ) -> dict:
        normalized_status = str(status or "").strip().lower()
        if normalized_status not in REVIEW_STATUSES - {"pending"}:
            raise ReviewServiceError(400, "status must be accepted, edited, or ignored", "ValidationError")

        items = self._items_with_state(job_id, workbook_path)
        selected = next((item for item in items if item["item_id"] == item_id), None)
        if not selected:
            raise ReviewServiceError(404, "review item not found", "ReviewItemNotFound")

        if normalized_status == "edited" and reviewed_value is None:
            raise ReviewServiceError(400, "edited review item requires reviewed_value", "ValidationError")

        next_value = selected["current_value"] if normalized_status == "accepted" else selected["reviewed_value"]
        if normalized_status == "edited":
            next_value = "" if reviewed_value is None else str(reviewed_value)
        if normalized_status == "ignored":
            next_value = "" if reviewed_value is None else str(reviewed_value)

        updated_at = _now_iso()
        state = self._read_state(job_id)
        state.setdefault("items", {})[item_id] = {
            "status": normalized_status,
            "reviewed_value": next_value,
            "comment": "" if comment is None else str(comment),
            "updated_at": updated_at,
        }
        state["updated_at"] = updated_at
        self._write_state(job_id, state)

        selected.update(state["items"][item_id])
        return self._public_item(selected)

    def apply_reviews(self, job_id: str, workbook_path: str | Path) -> dict:
        workbook_path = self._safe_source_workbook_path(workbook_path)
        items = self._items_with_state(job_id, workbook_path)
        workbook = load_workbook(workbook_path)
        try:
            if REVIEW_SHEET not in workbook.sheetnames:
                raise ReviewServiceError(404, "review sheet not found", "ReviewSheetNotFound")
            review_sheet = workbook[REVIEW_SHEET]
            review_headers = self._ensure_review_columns(review_sheet)
            result_sheet = workbook[RESULT_SHEET] if RESULT_SHEET in workbook.sheetnames else None
            result_headers = self._headers(result_sheet) if result_sheet is not None else []

            applied_items = [item for item in items if item["status"] in {"accepted", "edited", "ignored"}]
            for item in items:
                row_number = int(item["_sheet_row"])
                self._set_cell(review_sheet, review_headers, row_number, "review_status", item["status"])
                self._set_cell(review_sheet, review_headers, row_number, "reviewed_value", item["reviewed_value"])
                self._set_cell(review_sheet, review_headers, row_number, "review_comment", item["comment"])
                self._set_cell(review_sheet, review_headers, row_number, "updated_at", item["updated_at"])
                self._set_cell(review_sheet, review_headers, row_number, "item_id", item["item_id"])

            if result_sheet is not None:
                self._apply_edited_values(result_sheet, result_headers, applied_items)

            self._write_review_log(workbook, applied_items)
            reviewed_path = self._new_reviewed_workbook_path(job_id, workbook_path)
            workbook.save(reviewed_path)
        finally:
            workbook.close()

        state = self._read_state(job_id)
        updated_at = _now_iso()
        state["reviewed_workbook_file"] = reviewed_path.name
        state["updated_at"] = updated_at
        self._write_state(job_id, state)

        return {
            "job_id": job_id,
            "total": len(items),
            "applied_count": len(applied_items),
            "accepted_count": sum(1 for item in applied_items if item["status"] == "accepted"),
            "edited_count": sum(1 for item in applied_items if item["status"] == "edited"),
            "ignored_count": sum(1 for item in applied_items if item["status"] == "ignored"),
            "reviewed_workbook_url": f"/api/jobs/{job_id}/reviewed-workbook",
            "reviewed_file_name": reviewed_path.name,
            "updated_at": updated_at,
        }

    def reviewed_workbook_path(self, job_id: str) -> Path:
        state = self._read_state(job_id)
        file_name = str(state.get("reviewed_workbook_file") or "")
        if not file_name:
            raise ReviewServiceError(404, "reviewed workbook not ready", "ReviewedWorkbookNotReady")
        path = (self._job_dir(job_id) / file_name).resolve()
        if not self._is_relative_to(path, self._job_dir(job_id)) or path.suffix.lower() != ".xlsx" or not path.exists():
            raise ReviewServiceError(404, "reviewed workbook not found", "ReviewedWorkbookNotReady")
        return path

    def _items_with_state(self, job_id: str, workbook_path: str | Path) -> list[dict[str, Any]]:
        items = self._read_review_items(workbook_path)
        states = dict(self._read_state(job_id).get("items") or {})
        merged: list[dict[str, Any]] = []
        for item in items:
            state = states.get(item["item_id"])
            if isinstance(state, dict):
                item["status"] = self._safe_status(state.get("status"), item["status"])
                item["reviewed_value"] = _text(state.get("reviewed_value"))
                item["comment"] = _text(state.get("comment"))
                item["updated_at"] = _text(state.get("updated_at"))
            merged.append(item)
        return merged

    def _read_review_items(self, workbook_path: str | Path) -> list[dict[str, Any]]:
        workbook_path = self._safe_source_workbook_path(workbook_path)
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        try:
            if REVIEW_SHEET not in workbook.sheetnames:
                return []
            sheet = workbook[REVIEW_SHEET]
            headers = self._headers(sheet)
            items: list[dict[str, Any]] = []
            for sheet_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if not any(value not in (None, "") for value in row):
                    continue
                row_data = {headers[index]: row[index] if index < len(row) else "" for index in range(len(headers))}
                item_id = self._item_id(row_data, sheet_row)
                items.append(
                    {
                        "item_id": item_id,
                        "row_index": _safe_int(row_data.get("row_index")),
                        "field": _text(row_data.get("field")),
                        "current_value": _text(row_data.get("current_value")),
                        "reviewed_value": _text(row_data.get("reviewed_value")),
                        "confidence": _safe_float(row_data.get("confidence")),
                        "reason": _text(row_data.get("reason")),
                        "evidence": _text(row_data.get("evidence")),
                        "status": self._safe_status(row_data.get("review_status"), "pending"),
                        "comment": _text(row_data.get("review_comment")),
                        "updated_at": _text(row_data.get("updated_at")),
                        "source_id": _text(row_data.get("source_id")),
                        "info_id": _text(row_data.get("info_id")),
                        "title": _text(row_data.get("title")),
                        "source_url": _text(row_data.get("source_url")),
                        "evidence_id": _text(row_data.get("evidence_id")),
                        "attempt": _text(row_data.get("attempt")),
                        "suggested_action": _text(row_data.get("suggested_action")),
                        "_sheet_row": sheet_row,
                    }
                )
            return items
        finally:
            workbook.close()

    def _item_id(self, row_data: dict[str, Any], sheet_row: int) -> str:
        payload = {
            "source_id": _text(row_data.get("source_id")),
            "info_id": _text(row_data.get("info_id")),
            "row_index": _text(row_data.get("row_index")),
            "field": _text(row_data.get("field")),
            "evidence_id": _text(row_data.get("evidence_id")),
            "sheet_row": sheet_row,
        }
        digest = hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()
        return f"rev_{digest[:16]}"

    def _safe_source_workbook_path(self, workbook_path: str | Path) -> Path:
        path = Path(workbook_path).resolve()
        if not self._is_relative_to(path, self.output_root) or path.suffix.lower() != ".xlsx" or not path.exists():
            raise ReviewServiceError(404, "job workbook not found", "JobWorkbookNotFound")
        return path

    def _job_dir(self, job_id: str) -> Path:
        if not JOB_ID_RE.fullmatch(str(job_id or "")):
            raise ReviewServiceError(400, "invalid job_id", "ValidationError")
        path = (self.root / job_id).resolve()
        if not self._is_relative_to(path, self.root):
            raise ReviewServiceError(400, "invalid job_id", "ValidationError")
        return path

    def _state_path(self, job_id: str) -> Path:
        return self._job_dir(job_id) / STATE_FILENAME

    def _read_state(self, job_id: str) -> dict[str, Any]:
        path = self._state_path(job_id)
        if not path.exists():
            return {"job_id": job_id, "items": {}, "reviewed_workbook_file": "", "updated_at": ""}
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except Exception as exc:
            raise ReviewServiceError(500, f"review state read failed: {exc}", "ReviewStateReadError") from exc
        if not isinstance(payload, dict) or payload.get("job_id") != job_id:
            raise ReviewServiceError(500, "review state invalid", "ReviewStateReadError")
        payload.setdefault("items", {})
        payload.setdefault("reviewed_workbook_file", "")
        payload.setdefault("updated_at", "")
        return payload

    def _write_state(self, job_id: str, state: dict[str, Any]) -> None:
        path = self._state_path(job_id)
        ensure_dir(path.parent)
        payload = {
            "job_id": job_id,
            "items": state.get("items") if isinstance(state.get("items"), dict) else {},
            "reviewed_workbook_file": Path(str(state.get("reviewed_workbook_file") or "")).name,
            "updated_at": _text(state.get("updated_at")),
        }
        temp_path = path.with_suffix(".json.tmp")
        temp_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
        temp_path.replace(path)

    def _new_reviewed_workbook_path(self, job_id: str, workbook_path: Path) -> Path:
        directory = self._job_dir(job_id)
        ensure_dir(directory)
        safe_stem = re.sub(r"[^A-Za-z0-9._-]+", "_", workbook_path.stem).strip("._") or "workbook"
        path = (directory / f"{safe_stem}_reviewed_{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:8]}.xlsx").resolve()
        if not self._is_relative_to(path, directory):
            raise ReviewServiceError(400, "reviewed workbook path invalid", "ValidationError")
        return path

    def _ensure_review_columns(self, worksheet) -> dict[str, int]:
        headers = self._headers(worksheet)
        for header in ["item_id", "review_status", "reviewed_value", "review_comment", "updated_at"]:
            if header not in headers:
                worksheet.cell(row=1, column=len(headers) + 1).value = header
                headers.append(header)
        for index, header in enumerate(headers, start=1):
            cell = worksheet.cell(row=1, column=index)
            cell.value = header
            cell.font = Font(bold=True)
            if cell.fill is None or cell.fill.fill_type is None:
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
        return {header: index + 1 for index, header in enumerate(headers)}

    def _set_cell(self, worksheet, headers: dict[str, int], row_number: int, header: str, value: Any) -> None:
        column = headers.get(header)
        if not column:
            return
        worksheet.cell(row=row_number, column=column).value = sanitize_excel_value(value)

    def _apply_edited_values(self, worksheet, headers: list[str], items: list[dict[str, Any]]) -> None:
        header_to_column = {str(header): index + 1 for index, header in enumerate(headers)}
        for item in items:
            if item["status"] != "edited":
                continue
            field = item["field"]
            row_index = _safe_int(item["row_index"])
            column = header_to_column.get(field)
            if not column or row_index < 1:
                continue
            excel_row = row_index + 1
            if excel_row > worksheet.max_row:
                continue
            worksheet.cell(row=excel_row, column=column).value = sanitize_excel_value(item["reviewed_value"])

    def _write_review_log(self, workbook, items: list[dict[str, Any]]) -> None:
        if REVIEW_LOG_SHEET in workbook.sheetnames:
            workbook.remove(workbook[REVIEW_LOG_SHEET])
        worksheet = workbook.create_sheet(REVIEW_LOG_SHEET)
        headers = [
            "item_id",
            "row_index",
            "field",
            "status",
            "current_value",
            "reviewed_value",
            "comment",
            "updated_at",
            "source_id",
            "info_id",
        ]
        for column_index, header in enumerate(headers, start=1):
            cell = worksheet.cell(row=1, column=column_index)
            cell.value = header
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        for item in items:
            worksheet.append([sanitize_excel_value(item.get(header, "")) for header in headers])
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{worksheet.max_row}"
        for column_index, header in enumerate(headers, start=1):
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(len(header) + 8, 12), 42)

    def _headers(self, worksheet) -> list[str]:
        if worksheet is None or worksheet.max_row < 1:
            return []
        return ["" if cell.value is None else str(cell.value) for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]

    def _safe_status(self, value: Any, default: str) -> str:
        status = str(value or "").strip().lower()
        return status if status in REVIEW_STATUSES else default

    def _public_item(self, item: dict[str, Any]) -> dict:
        public = {
            "item_id": item["item_id"],
            "row_index": _safe_int(item.get("row_index")),
            "field": self._safe_text(item.get("field")),
            "current_value": self._safe_text(item.get("current_value")),
            "reviewed_value": self._safe_text(item.get("reviewed_value")),
            "confidence": _safe_float(item.get("confidence")),
            "reason": self._safe_text(item.get("reason")),
            "evidence": self._safe_text(item.get("evidence")),
            "status": self._safe_status(item.get("status"), "pending"),
            "comment": self._safe_text(item.get("comment")),
            "updated_at": self._safe_text(item.get("updated_at")),
        }
        for optional in ("source_id", "info_id", "title", "source_url", "evidence_id", "attempt", "suggested_action"):
            if optional == "source_url":
                public[optional] = self._safe_url(item.get(optional))
            else:
                public[optional] = self._safe_text(item.get(optional))
        return public

    def _safe_text(self, value: Any) -> str:
        text = _text(value)
        text = mask_sensitive_text(text)
        text = WINDOWS_PATH_RE.sub("[path]", text)
        return POSIX_PATH_RE.sub("[path]", text)

    def _safe_url(self, value: Any) -> str:
        text = mask_sensitive_text(mask_url(_text(value)))
        return WINDOWS_PATH_RE.sub("[path]", text)

    def _is_relative_to(self, path: Path, root: Path) -> bool:
        try:
            path.relative_to(root)
            return True
        except ValueError:
            return False


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _text(value: Any) -> str:
    return "" if value is None else str(value)


def _safe_int(value: Any) -> int:
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _safe_float(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0
