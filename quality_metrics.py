import json
import math
import re
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Dict, List, Sequence

from openpyxl import load_workbook

from field_mapping import load_field_mapping, normalize_record_fields


CORE_FIELDS = ["病种名称", "报销比例", "起付标准", "补助限额", "人员类型", "保险类型"]
ALIGN_FIELDS = ["SourceURL", "source_url", "info_id", "相关资讯", "标题", "Title"]
EMPTY_NORMALIZED = {"", "--", "None", "none", "NULL", "null", "nan"}


@dataclass
class CellComparison:
    field: str
    generated: Any
    manual: Any
    similarity: float
    match_type: str


@dataclass
class RowComparison:
    generated_index: int | None
    manual_index: int | None
    similarity: float
    cells: List[CellComparison] = field(default_factory=list)


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text in EMPTY_NORMALIZED:
        return ""
    return re.sub(r"\s+", "", text)


def _normalize_percent(text: str) -> str:
    match = re.fullmatch(r"([0-9]+(?:\.[0-9]+)?)\s*[％%]", text)
    if not match:
        return ""
    number = float(match.group(1))
    return f"percent:{number:g}"


def _normalize_amount(text: str) -> str:
    match = re.fullmatch(r"([0-9]+(?:\.[0-9]+)?)\s*(亿元|万元|万|元)", text)
    if not match:
        return ""
    number = float(match.group(1))
    unit = match.group(2)
    if unit == "亿元":
        number *= 100000000
    elif unit in {"万元", "万"}:
        number *= 10000
    return f"amount:{number:g}"


def _normalize_date(text: str) -> str:
    match = re.search(r"(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})", text)
    if not match:
        return ""
    year, month, day = match.groups()
    return f"date:{int(year):04d}-{int(month):02d}-{int(day):02d}"


def normalize_for_compare(value: Any) -> str:
    text = _clean_text(value)
    if not text:
        return ""
    text = text.replace("％", "%").replace("（", "(").replace("）", ")")
    for normalizer in (_normalize_percent, _normalize_amount, _normalize_date):
        normalized = normalizer(text)
        if normalized:
            return normalized
    return text.lower()


def compare_values(generated: Any, manual: Any, fuzzy_threshold: float = 0.82) -> tuple[float, str]:
    generated_text = _clean_text(generated)
    manual_text = _clean_text(manual)
    if not generated_text and not manual_text:
        return 1.0, "empty"
    if not generated_text or not manual_text:
        return 0.0, "missing"
    if generated_text == manual_text:
        return 1.0, "strict"
    if normalize_for_compare(generated_text) == normalize_for_compare(manual_text):
        return 1.0, "normalized"
    ratio = SequenceMatcher(None, generated_text, manual_text).ratio()
    if ratio >= fuzzy_threshold:
        return ratio, "fuzzy"
    return ratio, "different"


def read_excel_rows(path: str | Path, sheet_name: str = "结果数据") -> List[Dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook[workbook.sheetnames[0]]
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value or "").strip() for value in rows[0]]
        mapping = load_field_mapping(None)
        result: List[Dict[str, Any]] = []
        for raw_row in rows[1:]:
            row = {headers[index]: raw_row[index] if index < len(raw_row) else "" for index in range(len(headers))}
            if any(_clean_text(value) for value in row.values()):
                result.append(normalize_record_fields(row, mapping))
        return result
    finally:
        workbook.close()


def _alignment_key(row: Dict[str, Any]) -> str:
    for field in ALIGN_FIELDS:
        value = row.get(field)
        cleaned = _clean_text(value)
        if cleaned:
            return f"{field}:{cleaned}"
    return ""


def _row_similarity(generated: Dict[str, Any], manual: Dict[str, Any], fields: Sequence[str]) -> float:
    if not fields:
        return 0.0
    scores = [compare_values(generated.get(field), manual.get(field))[0] for field in fields]
    return sum(scores) / len(scores)


def align_rows(generated_rows: Sequence[Dict[str, Any]], manual_rows: Sequence[Dict[str, Any]], fields: Sequence[str]) -> List[tuple[int | None, int | None]]:
    manual_by_key: Dict[str, int] = {}
    for index, row in enumerate(manual_rows):
        key = _alignment_key(row)
        if key and key not in manual_by_key:
            manual_by_key[key] = index

    used_manual: set[int] = set()
    pairs: List[tuple[int | None, int | None]] = []
    for generated_index, generated in enumerate(generated_rows):
        key = _alignment_key(generated)
        manual_index = manual_by_key.get(key) if key else None
        if manual_index is None or manual_index in used_manual:
            candidates = [
                (index, _row_similarity(generated, manual, fields))
                for index, manual in enumerate(manual_rows)
                if index not in used_manual
            ]
            manual_index = max(candidates, key=lambda item: item[1])[0] if candidates else None
        if manual_index is not None:
            used_manual.add(manual_index)
        pairs.append((generated_index, manual_index))

    for manual_index in range(len(manual_rows)):
        if manual_index not in used_manual:
            pairs.append((None, manual_index))
    return pairs


def compare_rows(
    generated_rows: Sequence[Dict[str, Any]],
    manual_rows: Sequence[Dict[str, Any]],
    fields: Sequence[str] | None = None,
    core_fields: Sequence[str] = CORE_FIELDS,
) -> Dict[str, Any]:
    mapping = load_field_mapping(None)
    fields = list(fields or mapping.headers)
    pairs = align_rows(generated_rows, manual_rows, fields)
    row_results: List[RowComparison] = []
    field_scores: Dict[str, List[float]] = {field: [] for field in fields}

    for generated_index, manual_index in pairs:
        generated = generated_rows[generated_index] if generated_index is not None else {}
        manual = manual_rows[manual_index] if manual_index is not None else {}
        cells: List[CellComparison] = []
        for field in fields:
            similarity, match_type = compare_values(generated.get(field), manual.get(field))
            field_scores[field].append(similarity)
            cells.append(CellComparison(field, generated.get(field, ""), manual.get(field, ""), similarity, match_type))
        row_similarity = sum(cell.similarity for cell in cells) / len(cells) if cells else 0.0
        row_results.append(RowComparison(generated_index, manual_index, row_similarity, cells))

    field_metrics = {
        field: {
            "accuracy": sum(1 for score in scores if math.isclose(score, 1.0)) / len(scores) if scores else 0.0,
            "similarity": sum(scores) / len(scores) if scores else 0.0,
            "count": len(scores),
        }
        for field, scores in field_scores.items()
    }
    all_scores = [score for scores in field_scores.values() for score in scores]
    core_scores = [score for field in core_fields for score in field_scores.get(field, [])]
    differences = [
        {
            "generated_row": None if row.generated_index is None else row.generated_index + 2,
            "manual_row": None if row.manual_index is None else row.manual_index + 2,
            "field": cell.field,
            "generated": cell.generated,
            "manual": cell.manual,
            "similarity": round(cell.similarity, 6),
            "match_type": cell.match_type,
        }
        for row in row_results
        for cell in row.cells
        if cell.similarity < 1.0
    ]
    return {
        "overall_similarity": sum(all_scores) / len(all_scores) if all_scores else 0.0,
        "core_field_similarity": sum(core_scores) / len(core_scores) if core_scores else 0.0,
        "field_metrics": field_metrics,
        "row_matches": [
            {
                "generated_row": None if row.generated_index is None else row.generated_index + 2,
                "manual_row": None if row.manual_index is None else row.manual_index + 2,
                "similarity": row.similarity,
            }
            for row in row_results
        ],
        "differences": differences,
        "row_count": {"generated": len(generated_rows), "manual": len(manual_rows), "aligned": len(pairs)},
    }


def dump_json_report(result: Dict[str, Any], output_json: str | Path) -> Path:
    path = Path(output_json)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return path
