import argparse
import json
from pathlib import Path
from typing import Any, Dict, List

import yaml
from openpyxl import Workbook, load_workbook

from config import PROJECT_ROOT
from quality_metrics import CORE_FIELDS, compare_rows, dump_json_report, read_excel_rows


DEFAULT_THRESHOLDS = {
    "low_confidence": 0.7,
    "minimum_overall_similarity": 0.85,
    "minimum_core_field_similarity": 0.85,
}


def _load_thresholds(path: str | Path | None = None) -> Dict[str, Any]:
    config_path = Path(path) if path else PROJECT_ROOT / "config" / "quality_thresholds.yml"
    if not config_path.is_absolute():
        config_path = PROJECT_ROOT / config_path
    if not config_path.exists():
        return dict(DEFAULT_THRESHOLDS)
    payload = yaml.safe_load(config_path.read_text(encoding="utf-8")) or {}
    merged = dict(DEFAULT_THRESHOLDS)
    merged.update(payload)
    return merged


def _empty(value: Any) -> bool:
    return str(value or "").strip().lower() in {"", "--", "none", "null", "nan"}


def _missing_core_fields(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    missing: List[Dict[str, Any]] = []
    for row_index, row in enumerate(rows, start=2):
        for field in CORE_FIELDS:
            if _empty(row.get(field)):
                missing.append({"row": row_index, "field": field})
    return missing


def _read_sheet_rows(path: str | Path, sheet_name: str) -> List[Dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            return []
        worksheet = workbook[sheet_name]
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value or "").strip() for value in rows[0]]
        return [
            {headers[index]: value for index, value in enumerate(raw_row) if index < len(headers)}
            for raw_row in rows[1:]
            if any(value not in (None, "") for value in raw_row)
        ]
    finally:
        workbook.close()


def _low_confidence_fields(path: str | Path, threshold: float) -> List[Dict[str, Any]]:
    rows = _read_sheet_rows(path, "字段置信度")
    low_rows: List[Dict[str, Any]] = []
    for row in rows:
        try:
            confidence = float(row.get("confidence") or 0)
        except Exception:
            confidence = 0.0
        if confidence < threshold:
            low_rows.append(row)
    return low_rows


def _write_table_sheet(workbook: Workbook, title: str, rows: List[Dict[str, Any]]) -> None:
    worksheet = workbook.create_sheet(title)
    headers = list(rows[0].keys()) if rows else ["message"]
    worksheet.append(headers)
    for row in rows:
        worksheet.append([json.dumps(row.get(header), ensure_ascii=False) if isinstance(row.get(header), (dict, list)) else row.get(header, "") for header in headers])


def _write_report_xlsx(result: Dict[str, Any], output_path: str | Path) -> Path:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "summary"
    summary.append(["metric", "value"])
    for key in [
        "overall_similarity",
        "core_field_similarity",
        "row_alignment_score",
        "missing_rows",
        "extra_rows",
        "conflict_count",
        "low_confidence_count",
        "passed",
    ]:
        summary.append([key, result.get(key, "")])
    field_rows = [{"field": field, **metrics} for field, metrics in result.get("field_metrics", {}).items()]
    _write_table_sheet(workbook, "field_metrics", field_rows)
    _write_table_sheet(workbook, "row_matches", result.get("row_matches", []))
    _write_table_sheet(workbook, "differences", result.get("differences", []))
    _write_table_sheet(workbook, "missing_core_fields", result.get("missing_core_fields", []))
    _write_table_sheet(workbook, "low_confidence_fields", result.get("low_confidence_fields", []))
    workbook.save(path)
    return path


def run_quality_eval(
    generated_xlsx: str | Path,
    manual_xlsx: str | Path,
    output_xlsx: str | Path,
    output_json: str | Path | None = None,
    thresholds_path: str | Path | None = None,
) -> Dict[str, Any]:
    thresholds = _load_thresholds(thresholds_path)
    generated_rows = read_excel_rows(generated_xlsx)
    manual_rows = read_excel_rows(manual_xlsx)
    result = compare_rows(generated_rows, manual_rows)
    row_matches = result.get("row_matches", [])
    row_alignment_score = sum(float(item.get("similarity") or 0) for item in row_matches) / len(row_matches) if row_matches else 0.0
    missing_rows = sum(1 for item in row_matches if item.get("generated_row") is None)
    extra_rows = sum(1 for item in row_matches if item.get("manual_row") is None)
    low_confidence = _low_confidence_fields(generated_xlsx, float(thresholds.get("low_confidence", 0.7)))
    conflict_count = len(_read_sheet_rows(generated_xlsx, "冲突证据"))

    result.update(
        {
            "row_alignment_score": row_alignment_score,
            "missing_core_fields": _missing_core_fields(generated_rows),
            "missing_rows": missing_rows,
            "extra_rows": extra_rows,
            "low_confidence_fields": low_confidence,
            "low_confidence_count": len(low_confidence),
            "conflict_count": conflict_count,
            "thresholds": thresholds,
        }
    )
    result["passed"] = (
        float(result.get("overall_similarity") or 0) >= float(thresholds.get("minimum_overall_similarity", 0.85))
        and float(result.get("core_field_similarity") or 0) >= float(thresholds.get("minimum_core_field_similarity", 0.85))
        and not result["missing_core_fields"]
    )
    _write_report_xlsx(result, output_xlsx)
    if output_json:
        dump_json_report(result, output_json)
    return result


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Evaluate generated extraction workbook against a manually reviewed workbook.",
        allow_abbrev=False,
    )
    parser.add_argument("positional_generated_xlsx", nargs="?")
    parser.add_argument("positional_manual_xlsx", nargs="?")
    parser.add_argument("--generated", dest="generated_xlsx")
    parser.add_argument("--manual", dest="manual_xlsx")
    parser.add_argument("--output-xlsx", "--output", dest="output_xlsx", default="reports/quality_eval.xlsx")
    parser.add_argument("--output-json", "--json-output", dest="output_json", default="reports/quality_eval.json")
    parser.add_argument("--thresholds", default="config/quality_thresholds.yml")
    args = parser.parse_args()
    generated_xlsx = args.generated_xlsx or args.positional_generated_xlsx
    manual_xlsx = args.manual_xlsx or args.positional_manual_xlsx
    if not generated_xlsx and not manual_xlsx:
        parser.error("generated and manual workbooks are required; pass --generated/--manual or positional generated_xlsx manual_xlsx.")
    if not generated_xlsx:
        parser.error("generated workbook is required; pass --generated or positional generated_xlsx.")
    if not manual_xlsx:
        parser.error("manual workbook is required; pass --manual or positional manual_xlsx.")
    result = run_quality_eval(generated_xlsx, manual_xlsx, args.output_xlsx, args.output_json, args.thresholds)
    print(json.dumps({"passed": result["passed"], "overall_similarity": result["overall_similarity"]}, ensure_ascii=False))
    return 0 if result["passed"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
