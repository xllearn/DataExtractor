import json
import logging
import re
import subprocess
import sys
import time
import uuid
import inspect
import shutil
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from email import policy
from email.parser import BytesParser
from pathlib import Path
from threading import Lock
from typing import Any, Callable, List, Optional
from urllib.parse import quote

from fastapi import FastAPI, HTTPException, Query, Request
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from pydantic import BaseModel

from config import PROJECT_ROOT, apply_llm_config, load_settings
from config_loader import ConfigError, load_db_config, parse_selected_ids, validate_db_config_ready
from db_reader import count_configured_records, fetch_configured_records
from image_ocr import get_ocr_status
from keyword_utils import expand_keyword_groups
from security_utils import mask_sensitive_text
from services.extraction_service import ExtractionCancelled
from services.extraction_service import ExtractionRequest as ServiceExtractionRequest
from services.extraction_service import ExtractionService
from services.job_store import JobStore, JobStoreError
from utils import ensure_dir


MAX_SELECTED = 50
MAX_UPLOAD_BYTES = 20 * 1024 * 1024
WEB_APP_VERSION = "20260701_image_table"
LOGGER = logging.getLogger(__name__)
SENSITIVE_NAME_RE = re.compile(r"\b(DATABASE_URL|DB_PASSWORD|LLM_API_KEY|Authorization|api_key|password|token|secret)\b", re.IGNORECASE)
SENSITIVE_QUERY_RE = re.compile(
    r"([?&](?:X-Amz-Signature|X-Amz-Credential|Signature|Expires|api_key|password|token|secret)=)[^&\s]+",
    re.IGNORECASE,
)
WINDOWS_ABS_PATH_RE = re.compile(r"[A-Za-z]:\\.*?(?=\s+(?:and|or)\s+|[,;\"'\]}]|$)", re.IGNORECASE)


class ExtractRequest(BaseModel):
    selected_ids: List[str]
    mode: str = "merge"
    no_ocr: bool = True
    no_llm: bool = False
    external_ocr_text: str = ""
    prompt_version: str = "v3"


class UploadedExtractRequest(BaseModel):
    upload_id: str
    mode: str = "merge"
    no_ocr: bool = True
    no_llm: bool = False
    external_ocr_text: str = ""
    prompt_version: str = "v3"


class ApiError(Exception):
    def __init__(self, status_code: int, detail: str, error_type: str):
        super().__init__(detail)
        self.status_code = status_code
        self.detail = mask_sensitive_text(detail)
        self.error_type = error_type


def _json_error(status_code: int, detail: object, error_type: str) -> JSONResponse:
    if isinstance(detail, dict):
        safe_detail: object = {str(key): mask_sensitive_text(str(value)) for key, value in detail.items()}
    else:
        safe_detail = mask_sensitive_text(str(detail or "请求失败"))
    return JSONResponse(status_code=status_code, content={"detail": safe_detail, "error_type": error_type})


def _display_config_path(config_path: str) -> str:
    path = Path(config_path)
    if not path.is_absolute():
        return str(path).replace("\\", "/")
    try:
        return str(path.resolve().relative_to(PROJECT_ROOT)).replace("\\", "/")
    except ValueError:
        return path.name


def _git_commit() -> str:
    try:
        completed = subprocess.run(
            ["git", "rev-parse", "--short=12", "HEAD"],
            cwd=PROJECT_ROOT,
            text=True,
            capture_output=True,
            timeout=5,
            check=True,
        )
        commit = completed.stdout.strip()
        return commit or "unknown"
    except Exception:
        return "unknown"


def _version_payload() -> dict:
    return {
        "project_root": str(PROJECT_ROOT),
        "cwd": str(Path.cwd()),
        "git_commit": _git_commit(),
        "web_app_version": WEB_APP_VERSION,
        "api_features": {
            "job_mode": True,
            "job_metadata_persistence": True,
            "pagination_count": True,
            "ocr_status": True,
        },
    }


def _database_status(config_path: str) -> tuple[bool, str]:
    try:
        db_config = load_db_config(config_path)
        if not db_config.exists:
            return False, "数据库配置文件不存在，请使用 --config 指定可用 db_config.yml"
        validate_db_config_ready(db_config)
        return True, "数据库配置可用"
    except ConfigError as exc:
        return False, mask_sensitive_text(str(exc))


def _ensure_database_ready(config_path: str) -> None:
    safe, reason = _database_status(config_path)
    if not safe:
        raise ApiError(
            400,
            f"数据库未配置，无法查询文章。请使用 --config 指定可用 db_config.yml：{reason}",
            "DatabaseNotConfigured",
        )


def _article_summary(record: dict) -> dict:
    return {
        "id": str(record.get("_source_id") or record.get("SourceURL") or ""),
        "info_id": str(record.get("info_id") or ""),
        "title": str(record.get("Title") or ""),
        "source_url": str(record.get("SourceURL") or ""),
        "audit_time": str(record.get("AuditTime") or ""),
        "region": str(record.get("areaname") or record.get("province") or ""),
        "insurance_type": str(record.get("insurancetypename") or ""),
    }


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _new_job_id() -> str:
    return f"job_{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:8]}"


def _job_not_found() -> ApiError:
    return ApiError(404, "任务不存在或已过期，请返回列表重新生成。", "JobNotFound")


def _mask_api_text(value: object) -> str:
    text = mask_sensitive_text(str(value or ""))
    text = SENSITIVE_QUERY_RE.sub(r"\1***", text)
    return SENSITIVE_NAME_RE.sub("[redacted]", text)


def _default_record_provider(config_path: str):
    def provider(keyword: str = "", selected_ids: Optional[List[str]] = None, limit: int = 50, offset: int = 0):
        config = load_db_config(config_path, require_ready=True)
        keyword_groups = expand_keyword_groups(keyword)
        active_keyword_groups = [] if selected_ids else keyword_groups
        items = fetch_configured_records(
            config,
            limit=limit,
            offset=offset,
            selected_ids=selected_ids or [],
            keyword_groups=active_keyword_groups,
            keyword_mode=config.query.keyword_mode,
        )
        total = count_configured_records(
            config,
            selected_ids=selected_ids or [],
            keyword_groups=active_keyword_groups,
            keyword_mode=config.query.keyword_mode,
        )
        return {"items": items, "total": total}

    return provider


def _default_extract_runner(config_path: str, field_config_path: str, llm_config_path: str, output_dir: Path, log_dir: Path):
    def runner(
        selected_ids: List[str],
        mode: str,
        no_ocr: bool,
        no_llm: bool,
        external_ocr_text: str = "",
        prompt_version: str = "v3",
    ) -> Path:
        ensure_dir(output_dir)
        ensure_dir(log_dir)
        started_at = time.time()
        command = [
            sys.executable,
            str(PROJECT_ROOT / "main.py"),
            "--config",
            config_path,
            "--field-config",
            field_config_path,
            "--llm-config",
            llm_config_path,
            "--selected-ids",
            ",".join(selected_ids),
            "--mode",
            mode,
            "--output-dir",
            str(output_dir),
            "--log-dir",
            str(log_dir),
            "--prompt-version",
            (prompt_version or "v3").strip() or "v3",
            "--save-intermediate",
        ]
        if no_ocr:
            command.append("--no-ocr")
        if no_llm:
            command.append("--no-llm")
        if external_ocr_text.strip():
            external_path = log_dir / f"external_ocr_{int(started_at)}_{uuid.uuid4().hex[:8]}.txt"
            external_path.write_text(external_ocr_text, encoding="utf-8")
            command.extend(["--ocr-text-file", str(external_path)])
        completed = subprocess.run(command, cwd=PROJECT_ROOT, text=True, capture_output=True, timeout=900)
        if completed.returncode != 0:
            message = mask_sensitive_text((completed.stderr or completed.stdout or "抽取失败").strip())
            raise RuntimeError(message[-1000:])
        candidates = [path for path in output_dir.glob("*.xlsx") if path.stat().st_mtime >= started_at - 1]
        if not candidates:
            raise RuntimeError("抽取完成但未找到输出 Excel")
        return max(candidates, key=lambda item: item.stat().st_mtime)

    return runner


def _safe_download_path(output_dir: Path, file_id: str) -> Path:
    if "/" in file_id or "\\" in file_id or ".." in file_id:
        raise HTTPException(status_code=400, detail="invalid file_id")
    root = output_dir.resolve()
    path = (root / file_id).resolve()
    if path.parent != root or path.suffix.lower() != ".xlsx" or not path.exists():
        raise HTTPException(status_code=404, detail="file not found")
    return path


def _provider_items_and_total(result: object) -> tuple[list, int]:
    if isinstance(result, dict):
        items = list(result.get("items") or [])
        total = int(result.get("total", len(items)) or 0)
        return items, total
    items = list(result or [])
    return items, len(items)


def _pagination_payload(items: list, total: int, limit: int, offset: int) -> dict:
    minimum_total = offset + len(items) if items else 0
    total = max(int(total or 0), minimum_total)
    page = offset // limit + 1
    total_pages = (total + limit - 1) // limit if total else 0
    return {
        "items": [_article_summary(row) for row in items],
        "total": total,
        "limit": limit,
        "offset": offset,
        "page": page,
        "page_size": limit,
        "total_pages": total_pages,
        "has_next": offset + limit < total,
        "has_prev": offset > 0,
    }


def _parse_bool_form(value: object, default: bool = False) -> bool:
    if value is None:
        return default
    normalized = str(value).strip().lower()
    if not normalized:
        return default
    return normalized in {"1", "true", "yes", "on"}


def _parse_multipart_upload(content_type: str, body: bytes) -> tuple[dict[str, str], str, bytes]:
    if "multipart/form-data" not in content_type.lower():
        raise ApiError(400, "上传请求必须使用 multipart/form-data", "ValidationError")
    message = BytesParser(policy=policy.default).parsebytes(
        f"Content-Type: {content_type}\r\nMIME-Version: 1.0\r\n\r\n".encode("utf-8") + body
    )
    if not message.is_multipart():
        raise ApiError(400, "上传表单格式不正确", "ValidationError")

    fields: dict[str, str] = {}
    filename = ""
    file_bytes = b""
    for part in message.iter_parts():
        disposition = str(part.get("content-disposition") or "")
        if "form-data" not in disposition:
            continue
        name = str(part.get_param("name", header="content-disposition") or "")
        payload = part.get_payload(decode=True) or b""
        part_filename = str(part.get_filename() or "")
        if name == "file" and part_filename:
            filename = part_filename
            file_bytes = payload
            continue
        if name:
            charset = part.get_content_charset() or "utf-8"
            fields[name] = payload.decode(charset, errors="replace")

    if not filename:
        raise ApiError(400, "缺少上传 Excel 文件", "ValidationError")
    if not file_bytes:
        raise ApiError(400, "上传 Excel 文件不能为空", "ValidationError")
    return fields, filename, file_bytes


def _new_upload_id() -> str:
    return f"upload_{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:8]}"


def _uploads_root(output_root: Path) -> Path:
    return (output_root / "uploads").resolve()


def _upload_metadata_path(output_root: Path, upload_id: str) -> Path:
    if not re.fullmatch(r"upload_\d{8}_\d{6}_[0-9a-f]{8}", str(upload_id or "")):
        raise ApiError(400, "upload_id 不合法", "ValidationError")
    root = _uploads_root(output_root)
    path = (root / upload_id / "upload.json").resolve()
    if not _is_relative_to(path, root):
        raise ApiError(400, "upload_id 不合法", "ValidationError")
    return path


def _public_upload_payload(payload: dict) -> dict:
    return {
        "upload_id": str(payload.get("upload_id") or ""),
        "file_name": str(payload.get("file_name") or ""),
        "size": int(payload.get("size") or 0),
        "created_at": str(payload.get("created_at") or ""),
        "used_by_job_id": str(payload.get("used_by_job_id") or ""),
    }


def _read_upload_metadata(output_root: Path, upload_id: str) -> dict:
    path = _upload_metadata_path(output_root, upload_id)
    if not path.exists():
        raise ApiError(404, "upload 不存在", "UploadNotFound")
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        raise ApiError(500, f"upload metadata 读取失败: {exc}", "UploadReadError") from exc
    if not isinstance(payload, dict) or payload.get("upload_id") != upload_id:
        raise ApiError(500, "upload metadata 不合法", "UploadReadError")
    return payload


def _write_upload_metadata(output_root: Path, payload: dict) -> dict:
    path = _upload_metadata_path(output_root, str(payload.get("upload_id") or ""))
    ensure_dir(path.parent)
    public = _public_upload_payload(payload)
    internal = dict(public)
    internal["path"] = str(payload.get("path") or "")
    temp_path = path.with_suffix(".json.tmp")
    temp_path.write_text(json.dumps(internal, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    temp_path.replace(path)
    return public


def _uploaded_xlsx_path_from_metadata(output_root: Path, upload_id: str) -> Path:
    payload = _read_upload_metadata(output_root, upload_id)
    path = Path(str(payload.get("path") or "")).resolve()
    root = _uploads_root(output_root)
    if not _is_relative_to(path, root) or path.suffix.lower() != ".xlsx" or not path.exists():
        raise ApiError(404, "upload 文件不存在", "UploadNotFound")
    return path


def _validate_uploaded_workbook(path: Path) -> None:
    title_headers = {"title", "标题"}
    content_headers = {"content", "context", "正文", "内容", "html"}
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ApiError(400, "上传文件不是有效的 .xlsx workbook", "ValidationError") from exc
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            try:
                header_row = next(rows)
            except StopIteration:
                continue
            headers = {str(value or "").strip().lower() for value in header_row if str(value or "").strip()}
            if headers & title_headers and headers & content_headers:
                return
        raise ApiError(400, "上传 Excel 至少需要可识别的标题列和正文列", "ValidationError")
    finally:
        workbook.close()


def _store_uploaded_xlsx(output_root: Path, upload_id: str, filename: str, file_bytes: bytes) -> dict:
    suffix = Path(str(filename).replace("\\", "/")).suffix.lower()
    if suffix != ".xlsx":
        raise ApiError(400, "只允许上传 .xlsx 文件", "ValidationError")
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise ApiError(413, f"上传文件超过 {MAX_UPLOAD_BYTES // (1024 * 1024)}MB 限制", "UploadTooLarge")

    upload_root = _uploads_root(output_root)
    target_dir = (upload_root / upload_id).resolve()
    if not _is_relative_to(target_dir, upload_root):
        raise ApiError(400, "上传路径不合法", "ValidationError")
    ensure_dir(target_dir)
    target = (target_dir / f"{upload_id}.xlsx").resolve()
    if target.parent != target_dir:
        raise ApiError(400, "上传路径不合法", "ValidationError")
    target.write_bytes(file_bytes)
    try:
        _validate_uploaded_workbook(target)
    except Exception:
        try:
            target.unlink()
        except OSError:
            pass
        raise
    return _write_upload_metadata(
        output_root,
        {
            "upload_id": upload_id,
            "file_name": f"{upload_id}.xlsx",
            "size": len(file_bytes),
            "created_at": _now_iso(),
            "used_by_job_id": "",
            "path": str(target),
        },
    )


def _call_runner(
    runner: Callable,
    selected: List[str],
    mode: str,
    no_ocr: bool,
    no_llm: bool,
    external_ocr_text: str,
    prompt_version: str,
) -> Path:
    args = [selected, mode, no_ocr, no_llm, external_ocr_text]
    try:
        parameters = inspect.signature(runner).parameters
        positional_count = sum(
            1
            for parameter in parameters.values()
            if parameter.kind in {inspect.Parameter.POSITIONAL_ONLY, inspect.Parameter.POSITIONAL_OR_KEYWORD}
        )
        supports_prompt_version = (
            any(parameter.kind == inspect.Parameter.VAR_POSITIONAL for parameter in parameters.values())
            or "prompt_version" in parameters
            or positional_count >= 6
        )
    except (TypeError, ValueError):
        supports_prompt_version = True
    if supports_prompt_version:
        args.append((prompt_version or "v3").strip() or "v3")
    try:
        return Path(runner(*args))
    except TypeError:
        return Path(runner(selected, mode, no_ocr, no_llm))


def _call_service(service: Any, request: ServiceExtractionRequest, progress_callback: Callable[[dict], None], cancel_checker: Callable[[], bool]):
    try:
        parameters = inspect.signature(service.run).parameters
    except (TypeError, ValueError):
        parameters = {}
    supports_callbacks = any(parameter.kind == inspect.Parameter.VAR_KEYWORD for parameter in parameters.values()) or {
        "progress_callback",
        "cancel_checker",
    } <= set(parameters)
    if supports_callbacks:
        return service.run(request, progress_callback=progress_callback, cancel_checker=cancel_checker)
    return service.run(request)


def _preview_workbook(path: Path) -> dict:
    if not path.exists():
        raise ApiError(404, "文件不存在", "FileNotFound")
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if "结果数据" not in workbook.sheetnames:
            raise ApiError(404, "Excel 中不存在结果数据 sheet", "PreviewSheetNotFound")
        sheet = workbook["结果数据"]
        headers = [cell.value or "" for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        rows = []
        for row in sheet.iter_rows(min_row=2, max_row=101, values_only=True):
            if not any(value not in (None, "") for value in row):
                continue
            rows.append(["" if value is None else value for value in row])
        return {"sheet": "结果数据", "headers": headers, "rows": rows, "row_count": len(rows)}
    finally:
        workbook.close()


def _is_relative_to(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root)
        return True
    except ValueError:
        return False


def _unique_paths(paths: list[Path]) -> list[Path]:
    seen = set()
    unique: list[Path] = []
    for path in paths:
        resolved = path.resolve()
        key = str(resolved).lower()
        if key in seen:
            continue
        seen.add(key)
        unique.append(resolved)
    return unique


def _job_bound_dirs(job: dict, output_root: Path, log_root: Path) -> list[Path]:
    job_id = str(job.get("job_id") or "")
    roots = [output_root.resolve(), log_root.resolve()]
    candidates = [
        log_root.resolve() / job_id,
        output_root.resolve() / "logs" / job_id,
        output_root.resolve() / job_id,
    ]
    log_dir_text = str(job.get("log_dir") or "").strip()
    if log_dir_text:
        log_dir = Path(log_dir_text)
        if not log_dir.is_absolute() and ".." not in log_dir.parts:
            for root in roots:
                candidate = (root / log_dir).resolve()
                if candidate.name == job_id and _is_relative_to(candidate, root):
                    candidates.append(candidate)
    return _unique_paths([candidate for candidate in candidates if candidate.name == job_id])


def _job_log_dir(job: dict, output_root: Path, log_root: Path) -> Path:
    candidates = _job_bound_dirs(job, output_root, log_root)
    for candidate in candidates:
        if (candidate / "run.log").exists():
            return candidate
    return candidates[0] if candidates else (log_root.resolve() / str(job.get("job_id") or "")).resolve()


def _summary_candidates(job: dict, output_root: Path, log_root: Path) -> list[Path]:
    job_dirs = _job_bound_dirs(job, output_root, log_root)
    candidates = [job_dir / "summary.json" for job_dir in job_dirs]
    summary_text = str(job.get("summary_path") or "").strip()
    if summary_text:
        summary_path = Path(summary_text)
        if not summary_path.is_absolute() and ".." not in summary_path.parts and summary_path.name == "summary.json":
            if len(summary_path.parts) == 1:
                candidates.extend(job_dir / "summary.json" for job_dir in job_dirs)
            else:
                allowed_parents = {str(job_dir.resolve()).lower() for job_dir in job_dirs}
                for root in (output_root.resolve(), log_root.resolve()):
                    candidate = (root / summary_path).resolve()
                    if str(candidate.parent).lower() in allowed_parents:
                        candidates.append(candidate)
    return _unique_paths(candidates)


def _job_summary_path(job: dict, output_root: Path, log_root: Path) -> Path:
    candidates = _summary_candidates(job, output_root, log_root)
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return candidates[0] if candidates else (log_root.resolve() / str(job.get("job_id") or "") / "summary.json")


def _safe_display_path(value: object, output_root: Path, log_root: Path) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    path = Path(text)
    if path.is_absolute():
        resolved = path.resolve()
        for root in (output_root.resolve(), log_root.resolve()):
            if _is_relative_to(resolved, root):
                return _mask_api_text(resolved.relative_to(root).as_posix())
        return _mask_api_text(path.name)
    if ".." in path.parts:
        return _mask_api_text(path.name)
    return _mask_api_text(path.as_posix())


def _safe_summary_payload(payload: object, output_root: Path, log_root: Path, key: str = "") -> object:
    if key == "output_excel_paths" and isinstance(payload, list):
        return [_safe_display_path(item, output_root, log_root) for item in payload]
    if isinstance(payload, dict):
        return {str(item_key): _safe_summary_payload(item_value, output_root, log_root, str(item_key)) for item_key, item_value in payload.items()}
    if isinstance(payload, list):
        return [_safe_summary_payload(item, output_root, log_root, key) for item in payload]
    if key in {"output_excel_path", "log_dir", "summary_path"} or key.endswith("_path"):
        return _safe_display_path(payload, output_root, log_root)
    if isinstance(payload, str):
        return _mask_summary_text(payload, output_root, log_root)
    return payload


def _mask_summary_text(value: str, output_root: Path, log_root: Path) -> str:
    text = WINDOWS_ABS_PATH_RE.sub("[path]", str(value or ""))
    for root in (output_root.resolve(), log_root.resolve(), PROJECT_ROOT.resolve()):
        for fragment in {str(root), root.as_posix()}:
            if fragment:
                text = text.replace(fragment, "[path]")
    return _mask_api_text(text)


def create_app(
    config_path: str = "config/db_config.yml",
    field_config_path: str = "config/field_mapping.yml",
    llm_config_path: str = "config/llm_config.yml",
    output_dir: str | Path = "outputs/web",
    log_dir: str | Path = "logs/web",
    record_provider: Optional[Callable[..., list]] = None,
    extract_runner: Optional[Callable[[List[str], str, bool, bool], Path]] = None,
    extraction_service: Optional[Any] = None,
    use_subprocess_runner: bool = False,
) -> FastAPI:
    output_root = Path(output_dir)
    if not output_root.is_absolute():
        output_root = PROJECT_ROOT / output_root
    log_root = Path(log_dir)
    if not log_root.is_absolute():
        log_root = PROJECT_ROOT / log_root
    uses_default_provider = record_provider is None
    provider = record_provider or _default_record_provider(config_path)
    runner = extract_runner or (
        _default_extract_runner(config_path, field_config_path, llm_config_path, output_root, log_root)
        if use_subprocess_runner
        else None
    )
    service = extraction_service
    if service is None and runner is None:
        service = ExtractionService(
            output_dir=output_root,
            log_dir=log_root,
            record_provider=provider,
            config_path=config_path,
            field_config_path=field_config_path,
            llm_config_path=llm_config_path,
        )
    requires_database_for_extract = record_provider is None and extract_runner is None and extraction_service is None
    upload_service = service
    if upload_service is None:
        upload_service = ExtractionService(
            output_dir=output_root,
            log_dir=log_root,
            record_provider=provider,
            config_path=config_path,
            field_config_path=field_config_path,
            llm_config_path=llm_config_path,
        )
    job_store = JobStore(output_root)
    executor = ThreadPoolExecutor(max_workers=2)
    jobs: dict[str, dict] = {}
    cancelled_jobs: set[str] = set()
    jobs_lock = Lock()

    app = FastAPI(title="DataExtractor API")
    web_dir = PROJECT_ROOT / "web"
    if web_dir.exists():
        app.mount("/web", StaticFiles(directory=web_dir), name="web")

    @app.exception_handler(ApiError)
    async def api_error_handler(_request: Request, exc: ApiError):
        LOGGER.warning("API error %s: %s", exc.error_type, exc.detail)
        return _json_error(exc.status_code, exc.detail, exc.error_type)

    @app.exception_handler(HTTPException)
    async def http_error_handler(_request: Request, exc: HTTPException):
        return _json_error(exc.status_code, exc.detail, "HTTPException")

    @app.exception_handler(Exception)
    async def unhandled_error_handler(_request: Request, exc: Exception):
        message = mask_sensitive_text(str(exc) or "Internal Server Error")
        LOGGER.error("Unhandled API error %s: %s", exc.__class__.__name__, message)
        return _json_error(500, message, exc.__class__.__name__)

    @app.get("/")
    def index():
        index_path = web_dir / "index.html"
        if not index_path.exists():
            raise HTTPException(status_code=404, detail="web index not found")
        return FileResponse(index_path)

    @app.get("/api/health")
    def health():
        return {"ok": True}

    @app.get("/api/version")
    def version():
        return _version_payload()

    @app.get("/api/config/status")
    def config_status():
        settings = load_settings()
        try:
            settings = apply_llm_config(settings, llm_config_path)
        except Exception:
            pass
        safe_to_query, reason = _database_status(config_path)
        ocr_status = get_ocr_status()
        return {
            "database_configured": safe_to_query,
            "llm_configured": bool(settings.llm_api_key),
            "ocr_available": bool(ocr_status.get("available")),
            "ocr_status_reason": mask_sensitive_text(str(ocr_status.get("reason") or "")),
            "ocr_install_hint": mask_sensitive_text(str(ocr_status.get("install_hint") or "")),
            "ocr_engine": str(ocr_status.get("engine") or "none"),
            "config_path": _display_config_path(config_path),
            "database_status_reason": reason,
            "safe_to_query": safe_to_query,
            "web_app_version": WEB_APP_VERSION,
        }

    @app.get("/api/articles")
    def articles(
        keyword: str = "",
        selected_ids: str = "",
        limit: int = Query(default=50, ge=1, le=200),
        offset: int = Query(default=0, ge=0),
    ):
        selected = parse_selected_ids(selected_ids)
        if uses_default_provider:
            _ensure_database_ready(config_path)
        result = provider(keyword=keyword, selected_ids=selected, limit=limit, offset=offset)
        rows, total = _provider_items_and_total(result)
        return _pagination_payload(rows, total, limit, offset)

    @app.post("/api/extract")
    def extract(request: ExtractRequest):
        selected = [str(item).strip() for item in request.selected_ids if str(item).strip()]
        if not selected:
            raise ApiError(400, "selected_ids 不能为空", "ValidationError")
        if len(selected) > MAX_SELECTED:
            raise ApiError(400, f"selected_ids 单次最多 {MAX_SELECTED} 条", "ValidationError")
        if request.mode not in {"single", "merge"}:
            raise ApiError(400, "mode 只能是 single 或 merge", "ValidationError")
        if requires_database_for_extract:
            _ensure_database_ready(config_path)
        now = _now_iso()
        job_id = _new_job_id()
        with jobs_lock:
            jobs[job_id] = job_store.create(
                job_id=job_id,
                status="queued",
                message="queued",
                selected_ids=selected,
                input_mode="configured-db",
                progress_current=0,
                progress_total=len(selected),
                ocr_available=get_ocr_status().get("available"),
                external_ocr_used=bool(request.external_ocr_text.strip()),
                created_at=now,
                updated_at=now,
            )

        def run_job() -> None:
            with jobs_lock:
                if job_id in cancelled_jobs:
                    jobs[job_id] = job_store.update(
                        job_id,
                        status="cancelled",
                        message="cancelled",
                        selected_ids=selected,
                        finished_at=_now_iso(),
                        updated_at=_now_iso(),
                    )
                    return
                started_at = _now_iso()
                jobs[job_id] = job_store.update(
                    job_id,
                    status="running",
                    message="running",
                    started_at=started_at,
                    updated_at=started_at,
                )
            try:
                service_result = None
                if runner is not None:
                    output_path = _call_runner(
                        runner,
                        selected,
                        request.mode,
                        request.no_ocr,
                        request.no_llm,
                        request.external_ocr_text,
                        request.prompt_version,
                    )
                else:
                    service_request = ServiceExtractionRequest(
                        selected_ids=selected,
                        mode=request.mode,
                        no_ocr=request.no_ocr,
                        no_llm=request.no_llm,
                        external_ocr_text=request.external_ocr_text,
                        prompt_version=request.prompt_version,
                        job_id=job_id,
                    )

                    def progress_callback(payload: dict) -> None:
                        with jobs_lock:
                            if job_id in cancelled_jobs:
                                return
                            jobs[job_id] = job_store.update(job_id, status="running", message="running", **payload)

                    def cancel_checker() -> bool:
                        return job_id in cancelled_jobs

                    service_result = _call_service(service, service_request, progress_callback, cancel_checker)
                    output_path = service_result.output_excel_path
                file_id = output_path.name
                with jobs_lock:
                    if job_id in cancelled_jobs:
                        jobs[job_id] = job_store.update(
                            job_id,
                            status="cancelled",
                            message="cancelled",
                            selected_ids=selected,
                            finished_at=_now_iso(),
                            updated_at=_now_iso(),
                        )
                        return
                    update_fields = {
                        "run_id": "",
                        "input_mode": "configured-db",
                        "selected_ids": selected,
                        "progress_current": len(selected),
                        "progress_total": len(selected),
                        "summary_path": "",
                        "log_dir": log_root,
                    }
                    if service_result is not None:
                        update_fields.update(
                            {
                                "run_id": service_result.run_id,
                                "input_mode": service_result.input_mode,
                                "selected_ids": service_result.selected_ids,
                                "keyword": service_result.keyword,
                                "progress_current": service_result.total_records,
                                "progress_total": service_result.total_records,
                                "summary_path": service_result.summary_path,
                                "log_dir": service_result.log_dir,
                            }
                        )
                    jobs[job_id] = job_store.update(
                        job_id,
                        status="success",
                        message="done",
                        file_id=file_id,
                        download_url=f"/api/download/{quote(file_id, safe='')}",
                        preview_url=f"/api/jobs/{job_id}/preview",
                        output_excel_path=output_path,
                        finished_at=_now_iso(),
                        updated_at=_now_iso(),
                        **update_fields,
                    )
            except ExtractionCancelled:
                with jobs_lock:
                    cancelled_jobs.add(job_id)
                    jobs[job_id] = job_store.update(
                        job_id,
                        status="cancelled",
                        message="cancelled",
                        selected_ids=selected,
                        finished_at=_now_iso(),
                        updated_at=_now_iso(),
                    )
            except Exception as exc:
                with jobs_lock:
                    if job_id in cancelled_jobs:
                        jobs[job_id] = job_store.update(
                            job_id,
                            status="cancelled",
                            message="cancelled",
                            selected_ids=selected,
                            finished_at=_now_iso(),
                            updated_at=_now_iso(),
                        )
                        return
                    jobs[job_id] = job_store.update(
                        job_id,
                        status="failed",
                        message="failed",
                        selected_ids=selected,
                        error=mask_sensitive_text(str(exc)),
                        finished_at=_now_iso(),
                        updated_at=_now_iso(),
                    )

        executor.submit(run_job)
        return {
            "job_id": job_id,
            "status": "queued",
            "status_url": f"/api/jobs/{job_id}",
            "result_page": f"/web/result.html?job_id={job_id}",
        }

    async def parse_upload_request(request: Request) -> tuple[dict[str, str], str, bytes]:
        try:
            content_length = int(request.headers.get("content-length") or "0")
        except ValueError:
            content_length = 0
        if content_length > MAX_UPLOAD_BYTES:
            raise ApiError(413, f"upload exceeds {MAX_UPLOAD_BYTES // (1024 * 1024)}MB limit", "UploadTooLarge")

        body = await request.body()
        if len(body) > MAX_UPLOAD_BYTES:
            raise ApiError(413, f"upload exceeds {MAX_UPLOAD_BYTES // (1024 * 1024)}MB limit", "UploadTooLarge")

        return _parse_multipart_upload(str(request.headers.get("content-type") or ""), body)

    def create_upload(filename: str, file_bytes: bytes) -> dict:
        upload_id = _new_upload_id()
        return _store_uploaded_xlsx(output_root, upload_id, filename, file_bytes)

    def mark_upload_used(upload_id: str, job_id: str) -> None:
        payload = _read_upload_metadata(output_root, upload_id)
        payload["used_by_job_id"] = job_id
        _write_upload_metadata(output_root, payload)

    def start_uploaded_job(upload_id: str, request_data: UploadedExtractRequest) -> dict:
        upload_path = _uploaded_xlsx_path_from_metadata(output_root, upload_id)
        mode = str(request_data.mode or "merge")
        if mode not in {"single", "merge"}:
            raise ApiError(400, "mode must be single or merge", "ValidationError")
        job_id = _new_job_id()
        mark_upload_used(upload_id, job_id)
        selected: List[str] = []

        now = _now_iso()
        with jobs_lock:
            jobs[job_id] = job_store.create(
                job_id=job_id,
                status="queued",
                message="queued",
                selected_ids=selected,
                input_mode="uploaded-xlsx",
                progress_current=0,
                progress_total=0,
                ocr_available=get_ocr_status().get("available"),
                external_ocr_used=bool(request_data.external_ocr_text.strip()),
                created_at=now,
                updated_at=now,
            )

        def run_upload_job() -> None:
            with jobs_lock:
                if job_id in cancelled_jobs:
                    jobs[job_id] = job_store.update(
                        job_id,
                        status="cancelled",
                        message="cancelled",
                        selected_ids=selected,
                        finished_at=_now_iso(),
                        updated_at=_now_iso(),
                    )
                    return
                started_at = _now_iso()
                jobs[job_id] = job_store.update(
                    job_id,
                    status="running",
                    message="running",
                    started_at=started_at,
                    updated_at=started_at,
                )
            try:
                service_request = ServiceExtractionRequest(
                    selected_ids=selected,
                    mode=mode,
                    no_ocr=request_data.no_ocr,
                    no_llm=request_data.no_llm,
                    external_ocr_text=request_data.external_ocr_text,
                    prompt_version=request_data.prompt_version,
                    input_xlsx=str(upload_path),
                    job_id=job_id,
                )

                def progress_callback(payload: dict) -> None:
                    with jobs_lock:
                        if job_id in cancelled_jobs:
                            return
                        jobs[job_id] = job_store.update(job_id, status="running", message="running", **payload)

                def cancel_checker() -> bool:
                    return job_id in cancelled_jobs

                service_result = _call_service(upload_service, service_request, progress_callback, cancel_checker)
                output_path = service_result.output_excel_path
                file_id = output_path.name
                with jobs_lock:
                    if job_id in cancelled_jobs:
                        jobs[job_id] = job_store.update(
                            job_id,
                            status="cancelled",
                            message="cancelled",
                            selected_ids=selected,
                            finished_at=_now_iso(),
                            updated_at=_now_iso(),
                        )
                        return
                    jobs[job_id] = job_store.update(
                        job_id,
                        status="success",
                        message="done",
                        file_id=file_id,
                        download_url=f"/api/download/{quote(file_id, safe='')}",
                        preview_url=f"/api/jobs/{job_id}/preview",
                        output_excel_path=output_path,
                        run_id=service_result.run_id,
                        input_mode=service_result.input_mode or "uploaded-xlsx",
                        selected_ids=service_result.selected_ids,
                        keyword=service_result.keyword,
                        progress_current=service_result.total_records,
                        progress_total=service_result.total_records,
                        summary_path=service_result.summary_path,
                        log_dir=service_result.log_dir,
                        finished_at=_now_iso(),
                        updated_at=_now_iso(),
                    )
            except ExtractionCancelled:
                with jobs_lock:
                    cancelled_jobs.add(job_id)
                    jobs[job_id] = job_store.update(
                        job_id,
                        status="cancelled",
                        message="cancelled",
                        selected_ids=selected,
                        finished_at=_now_iso(),
                        updated_at=_now_iso(),
                    )
            except Exception as exc:
                with jobs_lock:
                    if job_id in cancelled_jobs:
                        jobs[job_id] = job_store.update(
                            job_id,
                            status="cancelled",
                            message="cancelled",
                            selected_ids=selected,
                            finished_at=_now_iso(),
                            updated_at=_now_iso(),
                        )
                        return
                    jobs[job_id] = job_store.update(
                        job_id,
                        status="failed",
                        message="failed",
                        selected_ids=selected,
                        error=mask_sensitive_text(str(exc)),
                        finished_at=_now_iso(),
                        updated_at=_now_iso(),
                    )

        executor.submit(run_upload_job)
        return {
            "job_id": job_id,
            "status": "queued",
            "status_url": f"/api/jobs/{job_id}",
            "result_page": f"/web/result.html?job_id={job_id}",
        }

    @app.post("/api/uploads/excel")
    async def upload_excel(request: Request):
        _fields, filename, file_bytes = await parse_upload_request(request)
        return create_upload(filename, file_bytes)

    @app.get("/api/uploads")
    def upload_list(limit: int = Query(default=20, ge=1, le=100)):
        root = _uploads_root(output_root)
        ensure_dir(root)
        items: list[dict] = []
        for path in root.glob("upload_*/upload.json"):
            try:
                payload = json.loads(path.read_text(encoding="utf-8"))
            except Exception:
                continue
            if isinstance(payload, dict):
                items.append(_public_upload_payload(payload))
        items.sort(key=lambda item: str(item.get("created_at") or ""), reverse=True)
        return {"items": items[:limit], "total": len(items), "limit": limit}

    @app.delete("/api/uploads/{upload_id}")
    def upload_delete(upload_id: str):
        payload = _read_upload_metadata(output_root, upload_id)
        if payload.get("used_by_job_id"):
            raise ApiError(409, "upload 已被任务使用，不能删除", "UploadInUse")
        path = _upload_metadata_path(output_root, upload_id)
        directory = path.parent.resolve()
        root = _uploads_root(output_root)
        if not _is_relative_to(directory, root):
            raise ApiError(400, "upload_id 不合法", "ValidationError")
        shutil.rmtree(directory)
        return {"upload_id": upload_id, "deleted": True}

    @app.post("/api/extract/uploaded")
    def extract_uploaded(request: UploadedExtractRequest):
        return start_uploaded_job(request.upload_id, request)

    @app.post("/api/extract/upload")
    async def upload_extract(request: Request):
        fields, filename, file_bytes = await parse_upload_request(request)
        upload = create_upload(filename, file_bytes)
        request_data = UploadedExtractRequest(
            upload_id=upload["upload_id"],
            mode=str(fields.get("mode") or "merge"),
            no_ocr=_parse_bool_form(fields.get("no_ocr"), True),
            no_llm=_parse_bool_form(fields.get("no_llm"), False),
            external_ocr_text=str(fields.get("external_ocr_text") or ""),
            prompt_version=str(fields.get("prompt_version") or "v3"),
        )
        return start_uploaded_job(request_data.upload_id, request_data)

    @app.get("/api/jobs")
    def job_list(limit: int = Query(default=20, ge=1, le=100)):
        items = job_store.list(limit=limit)
        return {"items": items, "total": len(items), "limit": limit}

    def read_job_or_404(job_id: str) -> dict:
        with jobs_lock:
            job = dict(jobs.get(job_id) or {})
        if not job:
            try:
                job = job_store.read(job_id)
            except JobStoreError:
                raise _job_not_found()
        if not job:
            raise _job_not_found()
        return job

    @app.get("/api/jobs/{job_id}")
    def job_status(job_id: str):
        job = read_job_or_404(job_id)
        return job_store.public_payload(job)

    @app.get("/api/jobs/{job_id}/logs")
    def job_logs(job_id: str, tail: int = 200, level: str = ""):
        job = read_job_or_404(job_id)
        normalized_level = str(level or "").strip().lower()
        if normalized_level and normalized_level not in {"info", "warning", "error"}:
            raise ApiError(400, "level 只能是 info、warning 或 error", "ValidationError")
        tail = max(0, min(int(tail or 0), 500))
        log_path = _job_log_dir(job, output_root, log_root) / "run.log"
        lines: list[str] = []
        if log_path.exists():
            raw_lines = log_path.read_text(encoding="utf-8", errors="replace").splitlines()
            if normalized_level:
                marker = f"[{normalized_level.upper()}]"
                raw_lines = [line for line in raw_lines if marker in line.upper()]
            lines = [_mask_api_text(line) for line in raw_lines[-tail:]] if tail else []
        return {"job_id": job_id, "level": normalized_level, "tail": tail, "lines": lines}

    @app.get("/api/jobs/{job_id}/summary")
    def job_summary(job_id: str):
        job = read_job_or_404(job_id)
        summary_path = _job_summary_path(job, output_root, log_root)
        if not summary_path.exists():
            raise ApiError(404, "summary 尚未生成", "SummaryNotReady")
        try:
            payload = json.loads(summary_path.read_text(encoding="utf-8"))
        except Exception as exc:
            raise ApiError(500, f"summary 读取失败: {exc}", "SummaryReadError") from exc
        return _safe_summary_payload(payload, output_root, log_root)

    @app.get("/api/jobs/{job_id}/download")
    def job_download(job_id: str):
        job = read_job_or_404(job_id)
        if job.get("status") != "success":
            raise ApiError(400, "job 尚未成功，不能下载", "JobNotReady")
        file_id = str(job.get("file_id") or "")
        if not file_id and job.get("output_excel_path"):
            file_id = Path(str(job.get("output_excel_path"))).name
        path = _safe_download_path(output_root, file_id)
        return FileResponse(
            path,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=path.name,
        )

    @app.post("/api/jobs/{job_id}/cancel")
    def job_cancel(job_id: str):
        job = read_job_or_404(job_id)
        status = str(job.get("status") or "")
        if status == "cancelled":
            return job_store.public_payload(job)
        if status in {"success", "failed"}:
            raise ApiError(400, "job 已结束，不能取消", "JobAlreadyFinished")
        with jobs_lock:
            cancelled_jobs.add(job_id)
            jobs[job_id] = job_store.update(
                job_id,
                status="cancelled",
                message="cancelled",
                finished_at=_now_iso(),
                updated_at=_now_iso(),
            )
            return jobs[job_id]

    @app.get("/api/jobs/{job_id}/preview")
    def job_preview(job_id: str):
        job = read_job_or_404(job_id)
        if job.get("status") != "success":
            raise ApiError(400, "job 尚未完成，不能预览", "JobNotReady")
        file_id = str(job.get("file_id") or "")
        if not file_id and job.get("output_excel_path"):
            file_id = Path(str(job.get("output_excel_path"))).name
        if not file_id and job.get("output_path"):
            file_id = Path(str(job.get("output_path"))).name
        return _preview_workbook(_safe_download_path(output_root, file_id))

    @app.get("/api/download/{file_id:path}")
    def download(file_id: str):
        path = _safe_download_path(output_root, file_id)
        return FileResponse(
            path,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=path.name,
        )

    return app


app = create_app()


def main() -> int:
    import argparse
    import uvicorn

    parser = argparse.ArgumentParser(description="启动 DataExtractor Web API")
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=8000)
    parser.add_argument("--config", default="config/db_config.yml")
    parser.add_argument("--field-config", default="config/field_mapping.yml")
    parser.add_argument("--llm-config", default="config/llm_config.yml")
    args = parser.parse_args()
    version = _version_payload()
    print("DataExtractor API starting")
    print(f"cwd={version['cwd']}")
    print(f"project_root={version['project_root']}")
    print(f"git_commit={version['git_commit']}")
    print(f"config={args.config}")
    print(f"field_config={args.field_config}")
    print(f"llm_config={args.llm_config}")
    print(f"web_app_version={version['web_app_version']}")
    uvicorn.run(
        create_app(config_path=args.config, field_config_path=args.field_config, llm_config_path=args.llm_config),
        host=args.host,
        port=args.port,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
