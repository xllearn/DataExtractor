import argparse
from pathlib import Path
from typing import Any, Dict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from quality_metrics import CORE_FIELDS, compare_rows, dump_json_report, read_excel_rows


def _write_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[dict[str, Any]]) -> None:
    worksheet = workbook.create_sheet(title)
    for index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=index, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows:
        worksheet.append([row.get(header, "") for header in headers])
    worksheet.freeze_panes = "A2"


def write_compare_report(result: Dict[str, Any], output_xlsx: str | Path) -> Path:
    output_path = Path(output_xlsx)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_sheet(
        workbook,
        "总体指标",
        ["metric", "value"],
        [
            {"metric": "overall_similarity", "value": result["overall_similarity"]},
            {"metric": "core_field_similarity", "value": result["core_field_similarity"]},
            {"metric": "generated_rows", "value": result["row_count"]["generated"]},
            {"metric": "manual_rows", "value": result["row_count"]["manual"]},
            {"metric": "aligned_rows", "value": result["row_count"]["aligned"]},
            {"metric": "core_fields", "value": "、".join(CORE_FIELDS)},
        ],
    )
    _write_sheet(
        workbook,
        "字段指标",
        ["field", "accuracy", "similarity", "count"],
        [{"field": field, **metrics} for field, metrics in result["field_metrics"].items()],
    )
    _write_sheet(
        workbook,
        "行匹配",
        ["generated_row", "manual_row", "similarity"],
        result["row_matches"],
    )
    _write_sheet(
        workbook,
        "差异明细",
        ["generated_row", "manual_row", "field", "generated", "manual", "similarity", "match_type"],
        result["differences"],
    )
    workbook.save(output_path)
    return output_path


def compare_excel_files(
    generated: str | Path,
    manual: str | Path,
    output_xlsx: str | Path | None = None,
    output_json: str | Path | None = None,
) -> Dict[str, Any]:
    generated_rows = read_excel_rows(generated)
    manual_rows = read_excel_rows(manual)
    result = compare_rows(generated_rows, manual_rows)
    if output_xlsx:
        write_compare_report(result, output_xlsx)
    if output_json:
        dump_json_report(result, output_json)
    return result


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="对比系统生成 Excel 与人工标注 Excel 的字段级相似度")
    parser.add_argument("--generated", required=True, help="系统生成的 Excel")
    parser.add_argument("--manual", required=True, help="人工撰写的 Excel")
    parser.add_argument("--output", default="", help="输出 xlsx 报告路径")
    parser.add_argument("--json-output", default="", help="输出 JSON 报告路径")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    output_xlsx = args.output or None
    output_json = args.json_output or (str(Path(args.output).with_suffix(".json")) if args.output else None)
    result = compare_excel_files(args.generated, args.manual, output_xlsx=output_xlsx, output_json=output_json)
    print(f"overall_similarity={result['overall_similarity']:.6f}")
    print(f"core_field_similarity={result['core_field_similarity']:.6f}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
